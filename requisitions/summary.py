"""The candidate summary emailed to a department, and the department's approval of who moves on.

HR sends it from the requisition page, or sets it to go out every day or every week
(``python manage.py send_department_summaries``). It gives the counts at each stage, the profiles
of the shortlisted candidates and an Excel list of every applicant (no contact details), and asks
the department to approve who moves on to interviews. Their answers are recorded on each
application and passed back to the recruiter.
"""

from __future__ import annotations

import re
from io import BytesIO

from django.conf import settings
from django.db.models import Count, F, Q
from django.urls import reverse
from django.utils import timezone

from core.notify import hr_team, notify, send_email
from pipeline.models import Application, Stage
from pipeline.services import WorkflowError, log

from .models import Requisition

# Stages at which the department is asked to approve who moves on.
REVIEW_STAGES = [Stage.SHORTLISTED]
OPEN_STATUSES = [Application.Status.ACTIVE, Application.Status.ON_HOLD]
NOT_PROGRESSING = [Application.Status.REJECTED, Application.Status.WITHDRAWN, Application.Status.NOT_ENGAGED]
LATER_STAGES = [Stage.DOCUMENTS, Stage.OFFER, Stage.MEDICAL, Stage.ONBOARDING]
EMAIL_TABLE_LIMIT = 25
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def answered_now():
    """Applications whose department answer is for the stage they are at now."""
    return Q(department_review__gt="") & Q(department_review_stage=F("stage"))


def awaiting_review(requisition):
    """Shortlisted candidates the department has not yet approved or turned down."""
    return (requisition.applications.filter(stage__in=REVIEW_STAGES, status=Application.Status.ACTIVE)
            .exclude(answered_now()).select_related("candidate").order_by("-match_score", "applied_at"))


def stage_counts(requisition) -> list[tuple[str, int]]:
    apps = requisition.applications.all()
    by_stage = dict(apps.filter(status__in=OPEN_STATUSES).order_by().values_list("stage").annotate(n=Count("id")))
    positions = requisition.positions
    return [
        ("Applied in total", apps.count()),
        ("Waiting to be screened", by_stage.get(Stage.APPLIED, 0)),
        ("Shortlisted", by_stage.get(Stage.SHORTLISTED, 0)),
        ("In interviews", by_stage.get(Stage.INTERVIEW, 0)),
        ("Evaluation and decision", by_stage.get(Stage.DECISION, 0)),
        ("Documents, offer, medicals and onboarding", sum(by_stage.get(s, 0) for s in LATER_STAGES)),
        (f"Hired (of {positions} position{'s' if positions != 1 else ''})",
         apps.filter(status=Application.Status.HIRED).count()),
        ("Not progressing (rejected, withdrawn or not engaged)", apps.filter(status__in=NOT_PROGRESSING).count()),
    ]


def review_label(application) -> str:
    answer = application.current_department_review
    if answer:
        return application.get_department_review_display()
    if application.stage in REVIEW_STAGES and application.status == Application.Status.ACTIVE:
        return "Awaiting your approval"
    return ""


def profile_row(application) -> list[str]:
    c = application.candidate
    background = " · ".join(part for part in (
        c.qualification_short if c.highest_qualification else "", c.degree_class_short, c.course) if part)
    match = f"{float(application.match_score):g}% ({application.match_grade})" if application.matched_at else "Not scored"
    return [c.full_name, background or "—", c.experience_label, match, review_label(application)]


def recipients(requisition):
    """The department's people in the portal, plus the department mailbox."""
    department = requisition.department
    users = department.contacts()
    raiser = requisition.raised_by
    if raiser and raiser.is_active and raiser.is_department_user and raiser not in users:
        users.append(raiser)
    emails = sorted({u.email for u in users if u.email} | ({department.email} if department.email else set()))
    return users, emails


def summary_workbook(requisition) -> bytes:
    """Every applicant for the records, ranked by match. Names only: no email addresses or phone numbers."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "-", requisition.reference or "Candidates")[:31]
    ws.append(["Rank", "Candidate", "Qualification", "Class of degree", "Course", "Institution", "Years exp.",
               "Match %", "Grade", "Stage", "Status", "Department's answer", "Applied"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0033A1")
    apps = requisition.applications.select_related("candidate").order_by("-match_score", "applied_at")
    for rank, app in enumerate(apps, start=1):
        c = app.candidate
        ws.append([
            rank, c.full_name, c.get_highest_qualification_display(), c.degree_class_short, c.course, c.institution,
            float(c.years_experience) if c.years_experience is not None else None,
            float(app.match_score) if app.matched_at else None, app.match_grade, app.get_stage_display(),
            app.get_status_display(), review_label(app), timezone.localtime(app.applied_at).strftime("%Y-%m-%d"),
        ])
    for column, width in zip("ABCDEFGHIJKLM", [6, 28, 24, 14, 28, 30, 9, 9, 7, 22, 12, 24, 12]):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def build_summary(requisition, *, note: str = "") -> dict:
    """Subject, text, figures and table of the email (also used for the preview)."""
    awaiting = list(awaiting_review(requisition))
    shortlisted = list(requisition.applications.filter(stage__in=REVIEW_STAGES, status__in=OPEN_STATUSES)
                       .select_related("candidate").order_by("-match_score", "applied_at"))
    today = timezone.localdate()
    department = requisition.department
    lines = [f"Hello {department} team,", "",
             f"Here is where recruitment for {requisition.title} ({requisition.reference}) stands on "
             f"{today.day} {today:%B %Y}."]
    if note.strip():
        lines += ["", note.strip()]
    lines.append("")
    if awaiting:
        n = len(awaiting)
        lines.append(f"{n} shortlisted candidate{'s are' if n != 1 else ' is'} waiting for your approval. Please open "
                     "the portal and approve who should move on to interviews, or mark those who should not.")
    else:
        lines.append("No candidates are waiting for your approval right now.")
    lines += ["", "The full list of candidates is attached as an Excel file for your records."]
    rows = [profile_row(a) for a in shortlisted[:EMAIL_TABLE_LIMIT]]
    more = len(shortlisted) - len(rows)
    review_url = reverse("requisitions:review", args=[requisition.pk])
    return {
        "subject": f"Candidate summary: {requisition.title} ({requisition.reference})",
        "body": "\n".join(lines),
        "stats": stage_counts(requisition),
        "table": {
            "title": "Shortlisted candidates",
            "headers": ["Candidate", "Background", "Experience", "Match", "Your decision"],
            "rows": rows,
            "more": f"…and {more} more in the attached list." if more > 0 else "",
            "empty": "Nobody is on the shortlist at the moment.",
        },
        "awaiting": len(awaiting),
        "action_url": review_url if awaiting else requisition.get_absolute_url(),
        "action_label": "Review the shortlist" if awaiting else "Open the vacancy",
    }


def send_summary(requisition, *, actor=None, note: str = "") -> list[str]:
    """Email the summary to the department. Returns the addresses it went to."""
    users, to = recipients(requisition)
    if not to:
        raise WorkflowError(f"{requisition.department} has no email address yet. Add the department's email, or a "
                            "head of department with an email, under Admin → Departments / Users.")
    content = build_summary(requisition, note=note)
    cc = [actor.email if actor else ""] + [requisition.hr_owner.email if requisition.hr_owner else ""]
    filename = f"{re.sub(r'[^A-Za-z0-9_-]', '-', requisition.reference)}-candidates.xlsx"
    sent = send_email(to, content["subject"], content["body"], cc=cc,
                      attachments=[(filename, summary_workbook(requisition), XLSX)],
                      action_url=content["action_url"], action_label=content["action_label"],
                      extra={"stats": content["stats"], "table": content["table"]})
    if not sent:
        raise WorkflowError("The email could not be sent. Check the email (SMTP) settings and try again.")
    waiting = content["awaiting"]
    notify(users, content["subject"],
           f"{waiting} shortlisted candidate{'s' if waiting != 1 else ''} waiting for your approval." if waiting
           else "Where recruitment stands, stage by stage.",
           url=content["action_url"], email=False, exclude=actor)
    requisition.summary_sent_at = timezone.now()
    requisition.save(update_fields=["summary_sent_at"])
    who = f"by {actor.display_name}" if actor else "automatically"
    log(f"Candidate summary emailed {who} to {', '.join(to)}.", verb="summary", requisition=requisition, actor=actor)
    return to


def is_due(requisition, now=None) -> bool:
    """Whether the automatic summary should go out now (from SUMMARY_EMAIL_HOUR, weekdays only for daily ones)."""
    frequency = requisition.summary_frequency
    if frequency == Requisition.SummaryFrequency.OFF or requisition.status not in Requisition.ACTIVE_STATUSES:
        return False
    local = timezone.localtime(now or timezone.now())
    if local.hour < settings.RECRUITMENT["SUMMARY_EMAIL_HOUR"]:
        return False
    last = timezone.localtime(requisition.summary_sent_at).date() if requisition.summary_sent_at else None
    if frequency == Requisition.SummaryFrequency.DAILY:
        return local.weekday() < 5 and last != local.date()
    return last is None or (local.date() - last).days >= 7


def record_review(application, answer: str, *, actor, note: str = ""):
    """The department approves (or not) a shortlisted candidate moving on."""
    if application.stage not in REVIEW_STAGES or application.status != Application.Status.ACTIVE:
        raise WorkflowError("This candidate is no longer waiting for the department's approval.")
    if answer not in Application.DepartmentReview.values:
        raise WorkflowError("Unknown answer.")
    application.department_review = answer
    application.department_review_stage = application.stage
    application.department_review_note = note.strip()
    application.department_reviewed_by = actor
    application.department_reviewed_at = timezone.now()
    application.save(update_fields=["department_review", "department_review_stage", "department_review_note",
                                    "department_reviewed_by", "department_reviewed_at", "updated_at"])
    req = application.requisition
    what = "approved to move on to interviews" if answer == Application.DepartmentReview.APPROVED else \
        "not approved to move on"
    log(f"{req.department}: {what} by {actor.display_name}." + (f" {note.strip()}" if note.strip() else ""),
        verb="review", application=application, actor=actor)
    team = [req.hr_owner] if req.hr_owner else hr_team()
    notify(team, f"{application.candidate.full_name} {what}", note.strip(), url=application.get_absolute_url(),
           email=False, exclude=actor)
    if not awaiting_review(req).exists():
        answered = req.applications.filter(answered_now(), stage__in=REVIEW_STAGES, status=Application.Status.ACTIVE)
        approved = answered.filter(department_review=Application.DepartmentReview.APPROVED).count()
        declined = answered.count() - approved
        notify(team, f"{req.department} has reviewed the shortlist for {req.title}",
               f"{approved} approved to move on to interviews, {declined} not approved.",
               url=req.get_absolute_url() + "?stage=shortlisted", exclude=actor)
