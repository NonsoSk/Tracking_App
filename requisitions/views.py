from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from core.choices import EmploymentType
from core.models import Employee
from core.permissions import can_view_requisition, ensure, hr_required, requisitions_for, role_required, safe_next
from core.text import normalize
from pipeline.models import STAGE_ICONS, Application, Interview, Stage
from pipeline.services import WorkflowError, auto_shortlist, move_to_stage, set_status as set_app_status

from . import services
from .forms import AdvertForm, ApproveForm, NoteForm, PublishForm, RequisitionForm, ReturnForm
from .knowledge_base import all_terms
from .models import JobRole, Requisition, SkillTag
from .suggestions import draft_job_advert, suggest_requirements

REQUIREMENT_FIELDS = ["min_experience_years", "max_experience_years", "education_level", "min_degree_class",
                      "requires_nysc", "courses", "required_skills", "preferred_skills", "tools", "certifications",
                      "employment_type"]


@login_required
def requisition_list(request):
    qs = requisitions_for(request.user).annotate(
        n_apps=Count("applications", distinct=True),
        n_active=Count("applications", filter=Q(applications__status="active"), distinct=True),
        n_interview=Count("applications", filter=Q(applications__stage="interview"), distinct=True),
    )
    status = request.GET.get("status", "")
    if status == "active":
        qs = qs.filter(status__in=[Requisition.Status.OPEN, Requisition.Status.APPROVED, Requisition.Status.ON_HOLD])
    elif status == "pending":
        qs = qs.filter(status__in=[Requisition.Status.SUBMITTED, Requisition.Status.PENDING_MANAGEMENT,
                                   Requisition.Status.RETURNED, Requisition.Status.DRAFT])
    elif status:
        qs = qs.filter(status=status)
    q = request.GET.get("q", "").strip()
    if q:
        qs = qs.filter(Q(title__icontains=q) | Q(reference__icontains=q) | Q(department__name__icontains=q))
    return render(request, "requisitions/list.html", {
        "requisitions": qs, "status": status, "q": q, "statuses": Requisition.Status.choices,
    })


def _initial_from_query(request):
    initial = {}
    replace = request.GET.get("replace")
    if replace:
        employee = Employee.objects.filter(pk=replace).select_related("department").first()
        if employee:
            reason = {
                Employee.Status.RESIGNED: Requisition.Reason.RESIGNATION,
                Employee.Status.TERMINATED: Requisition.Reason.TERMINATION,
                Employee.Status.TRANSFERRED: Requisition.Reason.TRANSFER,
            }.get(employee.status, Requisition.Reason.RETIREMENT)
            initial.update({
                "replacing_employee": employee.pk, "reason": reason, "department": employee.department_id,
                "title": employee.job_title, "grade": employee.grade,
                "justification": f"Replacement for {employee.full_name} ({employee.staff_id})"
                                 + (f", retiring {employee.retirement_date:%d %b %Y}." if reason == "retirement" and employee.retirement_date else "."),
            })
            role = JobRole.objects.filter(title__iexact=employee.job_title).first()
            if role:
                initial["job_role"] = role.pk
    return initial


@role_required("can_raise_requisition")
def requisition_create(request):
    form = RequisitionForm(request.POST or None, user=request.user, initial=_initial_from_query(request))
    if request.method == "POST" and form.is_valid():
        requisition = form.save(commit=False)
        requisition.raised_by = request.user
        requisition.save()
        services.log(f"Requisition created by {request.user.display_name}.", verb="requisition",
                     requisition=requisition, actor=request.user)
        if "submit" in request.POST:
            services.submit(requisition, actor=request.user)
            messages.success(request, f"{requisition.reference} submitted to the hiring team.")
        else:
            messages.success(request, f"{requisition.reference} saved as draft.")
        return redirect(requisition)
    return render(request, "requisitions/form.html", {"form": form, "is_new": True})


@login_required
def requisition_edit(request, pk):
    requisition = get_object_or_404(Requisition, pk=pk)
    user = request.user
    ensure(user.is_hr or (requisition.is_editable_by_department and can_view_requisition(user, requisition)
                          and user.can_raise_requisition))
    before = {f: getattr(requisition, f) for f in REQUIREMENT_FIELDS}
    form = RequisitionForm(request.POST or None, instance=requisition, user=user)
    if request.method == "POST" and form.is_valid():
        requisition = form.save()
        if any(getattr(requisition, f) != before[f] for f in REQUIREMENT_FIELDS):
            services.requirements_changed(requisition)
        if "submit" in request.POST and requisition.is_editable_by_department:
            services.submit(requisition, actor=user)
            messages.success(request, f"{requisition.reference} submitted to the hiring team.")
        else:
            messages.success(request, "Requisition saved.")
        return redirect(requisition)
    return render(request, "requisitions/form.html", {"form": form, "requisition": requisition, "is_new": False})


@login_required
def requisition_detail(request, pk):
    requisition = get_object_or_404(Requisition.objects.select_related("department", "raised_by", "hr_owner",
                                                                       "replacing_employee", "job_role"), pk=pk)
    ensure(can_view_requisition(request.user, requisition))
    applications = requisition.applications.select_related("candidate").prefetch_related("interviews")
    stage_filter = request.GET.get("stage", "")
    status_filter = request.GET.get("status", "open")
    listed = applications
    if stage_filter:
        listed = listed.filter(stage=stage_filter)
    if status_filter == "open":
        listed = listed.filter(status__in=[Application.Status.ACTIVE, Application.Status.ON_HOLD])
    elif status_filter:
        listed = listed.filter(status=status_filter)
    listed = listed.order_by("-match_score", "applied_at")

    stages = [s for s in Stage]
    board = {s.value: [] for s in stages}
    for app in applications.filter(status__in=[Application.Status.ACTIVE, Application.Status.ON_HOLD, Application.Status.HIRED]):
        board[app.stage].append(app)
    for items in board.values():
        items.sort(key=lambda a: -a.match_score)
    if requisition.is_trainee_track:
        stages = [s for s in stages if s != Stage.OFFER and (s != Stage.MEDICAL or settings.RECRUITMENT["TRAINEE_REQUIRES_MEDICALS"])]

    counts = {row["stage"]: row["n"] for row in applications.filter(
        status__in=[Application.Status.ACTIVE, Application.Status.ON_HOLD]).values("stage").annotate(n=Count("id"))}
    grade_counts = {row["match_grade"]: row["n"] for row in applications.values("match_grade").annotate(n=Count("id"))}
    interviews = Interview.objects.filter(application__requisition=requisition).select_related(
        "application__candidate").prefetch_related("panel").order_by("-scheduled_at")
    return render(request, "requisitions/detail.html", {
        "req": requisition, "applications": listed, "board": [(s, board[s.value]) for s in stages],
        "stage_filter": stage_filter, "status_filter": status_filter, "counts": counts, "grade_counts": grade_counts,
        "total_apps": applications.count(), "interviews": interviews, "notes": requisition.notes.select_related("author"),
        "activity": requisition.activities.select_related("actor", "application__candidate")[:40],
        "note_form": NoteForm(), "return_form": ReturnForm(), "approve_form": ApproveForm(),
        "publish_form": PublishForm(initial={"closing_date": requisition.closing_date}),
        "stages": [s for s in Stage], "threshold": settings.RECRUITMENT["AUTO_SHORTLIST_THRESHOLD"],
        "careers_url": request.build_absolute_uri(f"/careers/jobs/{requisition.reference}/"),
        "tab": request.GET.get("tab", "candidates"), "STAGE_ICONS": STAGE_ICONS,
    })


@login_required
@require_POST
def requisition_action(request, pk, action):
    requisition = get_object_or_404(Requisition, pk=pk)
    user = request.user
    ensure(can_view_requisition(user, requisition))
    try:
        if action == "submit":
            ensure(user.can_raise_requisition)
            services.submit(requisition, actor=user)
            messages.success(request, "Submitted to the hiring team.")
        elif action == "approve":
            ensure(user.is_hr or user.is_management)
            form = ApproveForm(request.POST)
            form.is_valid()
            services.approve(requisition, actor=user, hr_owner=form.cleaned_data.get("hr_owner"))
            messages.success(request, f"Requisition is now: {requisition.get_status_display()}.")
        elif action == "return":
            ensure(user.is_hr or user.is_management)
            form = ReturnForm(request.POST)
            if form.is_valid():
                services.return_for_changes(requisition, actor=user, reason=form.cleaned_data["reason"])
                messages.info(request, "Returned to the department for changes.")
        elif action == "publish":
            ensure(user.is_hr)
            form = PublishForm(request.POST)
            form.is_valid()
            services.publish(requisition, actor=user, closing_date=form.cleaned_data.get("closing_date"))
            messages.success(request, "The vacancy is now live on the careers page.")
        elif action in {"hold", "close", "cancel"}:
            ensure(user.is_hr or (action == "cancel" and requisition.raised_by_id == user.pk))
            status = {"hold": Requisition.Status.ON_HOLD, "close": Requisition.Status.CLOSED,
                      "cancel": Requisition.Status.CANCELLED}[action]
            services.set_status(requisition, status, actor=user, note=request.POST.get("note", ""))
            messages.info(request, f"Requisition {requisition.get_status_display().lower()}.")
        else:
            messages.error(request, "Unknown action.")
    except WorkflowError as exc:
        messages.error(request, str(exc))
    return redirect(requisition)


@login_required
@require_POST
def add_note(request, pk):
    requisition = get_object_or_404(Requisition, pk=pk)
    ensure(can_view_requisition(request.user, requisition))
    form = NoteForm(request.POST)
    if form.is_valid():
        services.add_note(requisition, author=request.user, message=form.cleaned_data["message"])
        messages.success(request, "Message sent.")
    return redirect(requisition.get_absolute_url() + "?tab=messages#messages")


@hr_required
def edit_advert(request, pk):
    requisition = get_object_or_404(Requisition, pk=pk)
    if request.method == "POST" and "generate" in request.POST:
        requisition.advert = draft_job_advert(requisition)
        requisition.save(update_fields=["advert"])
        messages.info(request, "Draft advert generated" + (" with AI." if settings.ANTHROPIC_API_KEY else " from the requirements."))
        return redirect("requisitions:advert", pk=pk)
    form = AdvertForm(request.POST or None, instance=requisition)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Job advert saved.")
        return redirect(requisition)
    return render(request, "requisitions/advert.html", {"form": form, "req": requisition})


@hr_required
@require_POST
def rescore(request, pk):
    requisition = get_object_or_404(Requisition, pk=pk)
    services.requirements_changed(requisition)
    messages.success(request, "All applications re-scored against the current requirements.")
    return redirect(requisition)


@hr_required
@require_POST
def run_auto_shortlist(request, pk):
    requisition = get_object_or_404(Requisition, pk=pk)
    try:
        threshold = float(request.POST.get("threshold") or settings.RECRUITMENT["AUTO_SHORTLIST_THRESHOLD"])
        top_n = int(request.POST["top_n"]) if request.POST.get("top_n") else None
    except ValueError:
        messages.error(request, "Enter valid numbers.")
        return redirect(requisition)
    picked = auto_shortlist(requisition, actor=request.user, threshold=threshold, top_n=top_n)
    messages.success(request, f"Shortlisted {len(picked)} candidate(s)." if picked else
                     "No new candidates met the threshold.")
    return redirect(requisition.get_absolute_url() + "?stage=shortlisted")


@hr_required
@require_POST
def bulk_action(request, pk):
    requisition = get_object_or_404(Requisition, pk=pk)
    ids = request.POST.getlist("selected")
    action = request.POST.get("bulk_action")
    applications = requisition.applications.filter(pk__in=ids)
    done = 0
    for application in applications:
        try:
            if action == "shortlist" and application.stage == Stage.APPLIED:
                move_to_stage(application, Stage.SHORTLISTED, actor=request.user, quiet=True)
            elif action == "reject":
                set_app_status(application, Application.Status.REJECTED, actor=request.user,
                               reason="Not shortlisted", email_candidate="email" in request.POST)
            elif action == "hold":
                set_app_status(application, Application.Status.ON_HOLD, actor=request.user, reason="Kept in reserve")
            else:
                continue
            done += 1
        except WorkflowError:
            continue
    if done and action == "shortlist":
        services.notify_department(requisition.department, f"{done} candidate(s) shortlisted for {requisition.title}",
                                   "Open the requisition to review the shortlist.", url=requisition.get_absolute_url())
    messages.success(request, f"Updated {done} application(s).")
    return redirect(safe_next(request, requisition.get_absolute_url(), request.META.get("HTTP_REFERER")))


@login_required
def export_applicants(request, pk):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    requisition = get_object_or_404(Requisition, pk=pk)
    ensure(can_view_requisition(request.user, requisition))
    wb = Workbook()
    ws = wb.active
    ws.title = requisition.reference
    headers = ["Rank", "Reference", "Name", "Email", "Phone", "Qualification", "Course", "Institution",
               "Years exp.", "Match %", "Grade", "Stage", "Status", "Source", "Panel score %", "Matched skills",
               "Missing skills", "Flags", "Applied"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B3D6E")
    for rank, app in enumerate(requisition.applications.select_related("candidate").order_by("-match_score"), start=1):
        c = app.candidate
        comps = app.match_breakdown.get("components", {})
        skills = comps.get("skills", {})
        summary = app.evaluation_summary()
        ws.append([
            rank, app.reference, c.full_name, c.email, c.phone, c.get_highest_qualification_display(), c.course,
            c.institution, float(c.years_experience), float(app.match_score), app.match_grade,
            app.get_stage_display(), app.get_status_display(), c.get_source_display(),
            float(summary["percentage"]) if summary else None,
            ", ".join(m["term"] for m in skills.get("matched", [])), ", ".join(skills.get("missing", [])),
            "; ".join(app.match_breakdown.get("flags", [])), timezone.localtime(app.applied_at).strftime("%Y-%m-%d"),
        ])
    for column, width in zip("ABCDEFGHIJKLMNOPQRS", [6, 16, 24, 28, 16, 22, 24, 28, 9, 9, 7, 18, 11, 22, 12, 40, 40, 40, 12]):
        ws.column_dimensions[column].width = width
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{requisition.reference}-applicants.xlsx"'
    wb.save(response)
    return response


# ---------------------------------------------------------------------------
# JSON endpoints used by the requisition form
# ---------------------------------------------------------------------------
@login_required
def api_suggest(request):
    title = request.GET.get("title", "")
    department = request.GET.get("department", "")
    if department.isdigit():
        from core.models import Department

        department = str(Department.objects.filter(pk=department).first() or "")
    employment = dict(EmploymentType.choices).get(request.GET.get("employment_type", ""), "")
    data = suggest_requirements(title, department, employment, request.GET.get("context", ""))
    return JsonResponse(data)


def api_terms(request):
    """Auto-complete for tag inputs (also used by the public application form)."""
    kind = request.GET.get("kind") or None
    q = normalize(request.GET.get("q", ""))
    terms = set(all_terms(kind))
    tags = SkillTag.objects.all()
    if kind:
        tags = tags.filter(kind=kind)
    terms.update(tags.values_list("name", flat=True))
    result = sorted((t for t in terms if not q or q in normalize(t)), key=lambda t: (not normalize(t).startswith(q), t.lower()))
    return JsonResponse({"results": result[:25]})


@login_required
def api_role(request, pk):
    role = get_object_or_404(JobRole, pk=pk)
    return JsonResponse({
        "title": role.title, "grade": role.grade, "min_experience_years": role.default_min_experience,
        "education_level": role.default_education, "skills": role.default_skills, "tools": role.default_tools,
        "certifications": role.default_certifications, "courses": role.default_courses,
        "description": role.description,
    })
