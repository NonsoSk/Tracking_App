"""
CV / résumé reader.

``parse_cv(data, filename)`` turns an uploaded file (PDF, Word, text, or a
scanned image / scanned PDF) into structured candidate data:

1. Text is extracted locally (pypdfium2 for PDFs, python-docx for Word files,
   Tesseract OCR for scans when it is installed).
2. When Claude is configured, the original file is sent to it for structured
   extraction. Claude reads scanned PDFs and photos of hard-copy CVs directly,
   so OCR software is not required in that case.
3. A rule-based extractor always runs as well; it fills anything the AI left
   empty and is the only extractor when AI is off.
"""

from __future__ import annotations

import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime

from django.conf import settings
from django.utils import timezone

from core import ai
from core.choices import DegreeClass, EDUCATION_RANK, EducationLevel, NyscStatus
from core.text import clean_list, find_in_text, merge_unique, normalize

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {".txt", ".md", ".rtf"}
WORD_EXTENSIONS = {".docx", ".doc", ".odt"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".tif", ".tiff", ".bmp"}
SUPPORTED_EXTENSIONS = {".pdf"} | TEXT_EXTENSIONS | WORD_EXTENSIONS | IMAGE_EXTENSIONS
CLAUDE_IMAGE_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp",
                      ".gif": "image/gif"}


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------
@dataclass
class ExtractedText:
    text: str = ""
    method: str = "none"
    scanned: bool = False
    notes: list[str] = field(default_factory=list)


def ocr_available() -> bool:
    try:
        import pytesseract
    except ImportError:
        return False
    if settings.TESSERACT_CMD:
        pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD
        return os.path.exists(settings.TESSERACT_CMD)
    return shutil.which("tesseract") is not None


def _ocr_images(images) -> str:
    import pytesseract

    if settings.TESSERACT_CMD:
        pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD
    return "\n".join(pytesseract.image_to_string(img) for img in images)


def _pdf_extract(data: bytes) -> ExtractedText:
    import pypdfium2 as pdfium

    result = ExtractedText(method="pdf-text")
    pdf = pdfium.PdfDocument(data)
    try:
        pages = []
        for index in range(len(pdf)):
            page = pdf[index]
            textpage = page.get_textpage()
            pages.append(textpage.get_text_bounded())
            textpage.close()
            page.close()
        result.text = "\n".join(pages)
        if len(re.sub(r"\s", "", result.text)) < 80:
            result.scanned = True
            if ocr_available():
                images = []
                for index in range(min(len(pdf), 6)):
                    page = pdf[index]
                    images.append(page.render(scale=2.5).to_pil())
                    page.close()
                result.text = _ocr_images(images)
                result.method = "ocr"
            else:
                result.notes.append("Scanned PDF: no text layer and OCR is not installed.")
    finally:
        pdf.close()
    return result


def _docx_extract(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text.strip() and cell.text not in cells:
                    cells.append(cell.text)
            parts.append(" | ".join(cells))
    for section in document.sections:
        parts = [p.text for p in section.header.paragraphs] + parts
    return "\n".join(parts)


def _office_convert_to_text(data: bytes, extension: str) -> str:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return ""
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, f"cv{extension}")
        with open(src, "wb") as fh:
            fh.write(data)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "txt:Text", "--outdir", tmp, src],
                check=True, capture_output=True, timeout=90,
            )
        except (subprocess.SubprocessError, OSError):
            return ""
        out = os.path.join(tmp, "cv.txt")
        if os.path.exists(out):
            with open(out, encoding="utf-8", errors="ignore") as fh:
                return fh.read()
    return ""


def _strip_rtf(text: str) -> str:
    text = re.sub(r"\\par[d]?", "\n", text)
    text = re.sub(r"\{\\\*[^{}]*\}", "", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    return re.sub(r"[{}]", "", text)


def extract_text(data: bytes, filename: str) -> ExtractedText:
    ext = os.path.splitext(filename or "")[1].lower()
    try:
        if ext == ".pdf":
            return _pdf_extract(data)
        if ext == ".docx":
            return ExtractedText(text=_docx_extract(data), method="docx")
        if ext in {".doc", ".odt"}:
            text = _office_convert_to_text(data, ext)
            notes = [] if text else ["Old Word format: install LibreOffice or save the CV as .docx/.pdf."]
            return ExtractedText(text=text, method="office" if text else "none", notes=notes)
        if ext in TEXT_EXTENSIONS:
            text = data.decode("utf-8", errors="ignore")
            return ExtractedText(text=_strip_rtf(text) if ext == ".rtf" else text, method="text")
        if ext in IMAGE_EXTENSIONS:
            result = ExtractedText(method="image", scanned=True)
            if ocr_available():
                from PIL import Image

                result.text = _ocr_images([Image.open(io.BytesIO(data))])
                result.method = "ocr"
            else:
                result.notes.append("Image CV: OCR is not installed.")
            return result
    except Exception as exc:  # corrupt files must never break an upload
        logger.warning("Could not read %s: %s", filename, exc)
        return ExtractedText(notes=[f"Could not read file: {exc}"])
    return ExtractedText(notes=[f"Unsupported file type {ext or '(none)'}"])


# ---------------------------------------------------------------------------
# Rule-based extraction
# ---------------------------------------------------------------------------
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+?234[\s-]?\(?0?\)?|\b0)[789][01]\d[\s-]?\d{3}[\s-]?\d{4}\b")
PHONE_FALLBACK_RE = re.compile(r"\+?\d[\d\s()-]{8,16}\d")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[A-Za-z0-9_-]+/?", re.I)

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
MONTH_PAT = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
DATE_PAT = rf"(?:{MONTH_PAT}\s*,?\s*|\d{{1,2}}\s*[/.-]\s*)?(?:19|20)\d{{2}}"
RANGE_RE = re.compile(
    rf"({DATE_PAT})\s*(?:-|–|—|to|till|until)\s*(present|current|date|now|till date|to date|{DATE_PAT})", re.I
)

SECTION_HEADINGS = {
    "summary": r"(?:professional |career |personal )?(?:summary|profile|objectives?|statement)|about me",
    "experience": r"(?:work |professional |employment |career |relevant |industrial )?(?:experience|history)"
                  r"|employment(?: record| history)?|work history|career history|positions held",
    "education": r"(?:education(?:al)?|academic)(?: background| qualifications?| history| records?)?"
                 r"|qualifications|institutions attended",
    "skills": r"(?:key |core |technical |professional |relevant )?(?:skills|competenc(?:ies|e)|expertise|strengths)"
              r"(?: (?:and|&) (?:competencies|abilities|tools))?|tools|software",
    "certifications": r"(?:professional )?(?:certifications?|certificates?|trainings?|licen[cs]es?|courses attended)"
                      r"(?: (?:and|&) (?:trainings?|certifications?|memberships?))?|professional (?:qualifications|development|memberships?)",
    "references": r"references?|referees?",
    "personal": r"personal (?:data|details|information)|bio ?data",
    "other": r"hobbies|interests|awards|achievements|languages|projects|publications|volunteer.*|extra.?curricular.*",
}
HEADING_RE = {k: re.compile(rf"^\s*(?:{v})\s*:?\s*$", re.I) for k, v in SECTION_HEADINGS.items()}

EDUCATION_PATTERNS = [
    (EducationLevel.PHD, r"\bph\.?\s?d\b|doctor of philosophy"),
    (EducationLevel.MSC, r"\bm\.?\s?sc\b|\bm\.?\s?eng\b|\bmba\b|\bm\.?\s?tech\b|\bm\.?\s?phil\b|\bllm\b|\bmasters?\b|master of"),
    (EducationLevel.PGD, r"\bpgd\b|post\s*-?graduate diploma"),
    (EducationLevel.BSC, r"\bb\.?\s?sc\b|\bb\.?\s?eng\b|\bb\.?\s?tech\b|\bbachelor|\bb\.a\.?|\bllb\b|\bb\.?\s?ed\b|\bmbbs\b|\bb\.?\s?pharm\b"),
    (EducationLevel.HND, r"\bhnd\b|higher national diploma"),
    (EducationLevel.OND, r"\bond\b|ordinary national diploma|\bnational diploma\b"),
    (EducationLevel.NCE, r"\bnce\b|national certificate (?:of|in) education"),
    (EducationLevel.SSCE, r"\bssce\b|\bwaec\b|\bneco\b|\bwassce\b|\bgce\b|senior secondary"),
]

DEGREE_CLASS_PATTERNS = [
    (DegreeClass.FIRST, r"first class|\bdistinction\b"),
    (DegreeClass.SECOND_UPPER, r"second class upper|second class \(upper|2\s*[:.]\s*1\b|upper credit|2nd class upper"),
    (DegreeClass.SECOND_LOWER, r"second class lower|second class \(lower|2\s*[:.]\s*2\b|lower credit|2nd class lower"),
    (DegreeClass.THIRD, r"third class|3rd class"),
]

COURSE_RE = re.compile(
    r"(?:b\.?\s?sc|b\.?\s?eng|b\.?\s?tech|m\.?\s?sc|m\.?\s?eng|m\.?\s?tech|hnd|ond|bachelor(?:'s)?(?: of \w+)?|"
    r"master(?:'s)?(?: of \w+)?|national diploma|degree)\.?\s*(?:\(hons?\.?\)|\(honours\)|hons\.?)?\s*"
    r"(?:degree)?\s*(?:in|,|-|–|:)?\s*\(?([A-Za-z][A-Za-z&/,' -]{2,70})",
    re.I,
)
INSTITUTION_RE = re.compile(r"(university|polytechnic|college|institute|school of|academy)", re.I)
LOCATIONS = ["Port Harcourt", "Eleme", "Onne", "Bonny", "Lagos", "Abuja", "Warri", "Kaduna", "Kano", "Enugu", "Owerri",
             "Uyo", "Calabar", "Benin City", "Ibadan", "Aba", "Yenagoa", "Asaba", "Abeokuta", "Ilorin", "Jos",
             "Akure", "Oyo", "Ogun", "Rivers State", "Delta State", "Bayelsa", "Akwa Ibom", "Anambra", "Imo"]
COMPANY_HINTS = re.compile(
    r"\b(ltd|limited|plc|nigeria|company|services|group|inc|llc|bank|industries|petroleum|oil|gas|energy|"
    r"fertili[sz]er|chemicals|consult\w*|engineering|corporation|hospital|university|agency|ministry)\b", re.I
)
SKIP_NAME_WORDS = re.compile(
    r"curriculum|vitae|resume|résumé|\bcv\b|profile|summary|email|e-mail|phone|tel|mobile|address|objective|"
    r"@|www|http|linkedin|\d", re.I
)
AMBIGUOUS_TERMS = {
    # "Plc" is a company suffix in Nigeria; require real PLC context.
    "plc": re.compile(r"\bplc\s*(programming|logic|control|systems?|troubleshoot|ladder|configuration)|programmable logic", re.I),
}


def _split_sections(lines: list[str]) -> dict[str, str]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in lines:
        stripped = line.strip().strip("•*-_=#").strip()
        heading = None
        if 0 < len(stripped) <= 60:
            for name, pattern in HEADING_RE.items():
                if pattern.match(stripped):
                    heading = name
                    break
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return {k: "\n".join(v) for k, v in sections.items()}


def _guess_name(lines: list[str]) -> tuple[str, str]:
    for line in lines[:12]:
        raw = line.strip()
        match = re.match(r"^(?:full\s+)?name\s*[:\-]\s*(.+)$", raw, re.I)
        if match:
            raw = match.group(1)
        elif SKIP_NAME_WORDS.search(raw) or raw.islower():
            continue
        words = [w for w in re.split(r"\s+", raw.replace(",", " ")) if w]
        if 2 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z][A-Za-z.'-]*", w) for w in words):
            if raw.isupper() or raw.islower():
                words = [w.capitalize() for w in words]
            words = [w for w in words if w.lower().rstrip(".") not in {"mr", "mrs", "miss", "ms", "engr", "dr"}]
            if len(words) >= 2:
                return words[0], " ".join(words[1:])
    return "", ""


def _parse_date_token(token: str):
    token = token.strip().lower()
    if token in {"present", "current", "date", "now", "till date", "to date"}:
        today = timezone.localdate()
        return today.year, today.month
    year_match = re.search(r"(19|20)\d{2}", token)
    if not year_match:
        return None
    year = int(year_match.group(0))
    month = 1
    name = re.match(r"([a-z]{3})", token)
    if name and name.group(1) in MONTHS:
        month = MONTHS[name.group(1)]
    else:
        num = re.match(r"(\d{1,2})\s*[/.-]", token)
        if num and 1 <= int(num.group(1)) <= 12:
            month = int(num.group(1))
    return year, month


def _experience_from_ranges(text: str) -> tuple[float, list[dict]]:
    intervals = []
    entries = []
    lines = text.splitlines()
    for idx, line in enumerate(lines):
        for match in RANGE_RE.finditer(line):
            start = _parse_date_token(match.group(1))
            end = _parse_date_token(match.group(2))
            if not start or not end:
                continue
            s = start[0] * 12 + start[1]
            e = end[0] * 12 + end[1]
            if e < s or e - s > 45 * 12:
                continue
            intervals.append((s, e))
            rest = (line[: match.start()] + line[match.end():]).strip(" |,-–—:()\t")
            context = rest or next((lines[j].strip() for j in range(idx - 1, max(idx - 3, -1), -1) if lines[j].strip()), "")
            title, company = _split_title_company(context)
            entries.append({"job_title": title, "company": company, "start": match.group(1).strip(),
                            "end": match.group(2).strip().title(), "description": ""})
    intervals.sort()
    merged: list[list[int]] = []
    for s, e in intervals:
        if merged and s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    months = sum(e - s for s, e in merged)
    return round(months / 12, 1), entries


def _split_title_company(text: str) -> tuple[str, str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return "", ""
    parts = [p.strip() for p in re.split(r"\s+at\s+|\s*[|,–—]\s*|\s+-\s+", text) if p.strip()]
    if len(parts) == 1:
        return ("", parts[0]) if COMPANY_HINTS.search(parts[0]) else (parts[0], "")
    company = next((p for p in parts if COMPANY_HINTS.search(p)), parts[1])
    title = next((p for p in parts if p != company), "")
    return title[:150], company[:150]


def _known_terms():
    """Catalogue terms grouped by kind (knowledge base + admin skill tags)."""
    from requisitions.knowledge_base import all_terms
    from requisitions.models import SkillTag

    catalogue = {kind: set(all_terms(kind)) for kind in ("skill", "tool", "certification", "course")}
    try:
        for name, kind in SkillTag.objects.values_list("name", "kind"):
            catalogue.setdefault(kind, set()).add(name)
    except Exception:  # table may not exist during migrations
        pass
    return catalogue


def detect_terms(text: str, kind: str, catalogue=None) -> list[str]:
    catalogue = catalogue or _known_terms()
    text_norm = normalize(text)
    found = []
    for term in sorted(catalogue.get(kind, ()), key=str.lower):
        key = normalize(term)
        if key in AMBIGUOUS_TERMS:
            if AMBIGUOUS_TERMS[key].search(text):
                found.append(term)
            continue
        if find_in_text(term, text_norm):
            found.append(term)
    return found


def _list_items(section_text: str, max_words: int) -> list[str]:
    items = []
    for line in section_text.splitlines():
        for item in re.split(r"[•;,|]|\s{3,}", line):
            item = re.sub(r"^[\s\-–*o·\d.)]+", "", item).strip()
            item = re.sub(r"\(?\b(19|20)\d{2}\b\)?", "", item).strip(" -–:")
            if 2 < len(item) <= 80 and len(item.split()) <= max_words:
                items.append(item)
    return items


def heuristic_extract(text: str) -> dict:
    data: dict = {}
    if not text.strip():
        return data
    lines = [l for l in text.splitlines()]
    non_empty = [l for l in lines if l.strip()]
    sections = _split_sections(lines)
    lower = text.lower()

    data["first_name"], data["last_name"] = _guess_name(non_empty)
    email = EMAIL_RE.search(text)
    data["email"] = email.group(0).lower() if email else ""
    phone = PHONE_RE.search(text) or PHONE_FALLBACK_RE.search(text)
    data["phone"] = re.sub(r"[^\d+]", "", phone.group(0)) if phone else ""
    linkedin = LINKEDIN_RE.search(text)
    data["linkedin_url"] = ("https://" + linkedin.group(0).split("://")[-1]) if linkedin else ""

    header = "\n".join(non_empty[:15]) + "\n" + sections.get("personal", "")
    data["location"] = next((loc for loc in LOCATIONS if re.search(rf"\b{re.escape(loc)}\b", header, re.I)), "")
    gender = re.search(r"\b(?:sex|gender)\s*[:\-]?\s*(male|female)\b", text, re.I)
    data["gender"] = gender.group(1).lower() if gender else ""
    dob = re.search(r"(?:date of birth|d\.?o\.?b\.?)\s*[:\-]?\s*([0-9A-Za-z ,/.-]{6,25})", text, re.I)
    data["date_of_birth"] = _parse_dob(dob.group(1)) if dob else None

    education = sections.get("education") or text
    edu_lower = education.lower()
    levels = [lvl for lvl, pattern in EDUCATION_PATTERNS if re.search(pattern, edu_lower)]
    data["highest_qualification"] = max(levels, key=lambda l: EDUCATION_RANK[l]) if levels else ""

    catalogue = _known_terms()
    known_courses = detect_terms(education, "course", catalogue)
    course = ""
    if known_courses:
        edu_norm = normalize(education)
        course = min(known_courses, key=lambda c: edu_norm.find(normalize(c)) if normalize(c) in edu_norm else 10**6)
    else:
        m = COURSE_RE.search(education)
        if m:
            course = re.split(r"\b(from|at|university|polytechnic|college|institute|\d)\b|[,(]", m.group(1), flags=re.I)[0]
            course = course.strip(" -–:,").title()
    data["course"] = course[:150]

    institution = next((l.strip() for l in education.splitlines() if INSTITUTION_RE.search(l)), "")
    institution = next((part for part in re.split(r"\s*[,|–—]\s*|\s+-\s+", institution) if INSTITUTION_RE.search(part)), institution)
    institution = re.sub(r"\(?\b(19|20)\d{2}\b\)?|[-–]\s*$", "", institution).strip(" ,-–|:")
    if len(institution) > 120:
        m = re.search(r"([A-Z][\w.'&-]*(?:\s+[\w.'&-]+){0,6}\s*(?:University|Polytechnic|College|Institute)[\w ,.'-]{0,40})", institution)
        institution = m.group(1) if m else institution[:120]
    data["institution"] = institution

    years = [int(y) for y in re.findall(r"\b((?:19|20)\d{2})\b", education) if int(y) <= timezone.localdate().year]
    data["graduation_year"] = max(years) if years else None
    data["degree_class"] = next((cls for cls, p in DEGREE_CLASS_PATTERNS if re.search(p, lower)), "")

    if re.search(r"\bnysc\b|national youth service", lower):
        if re.search(r"exempt", lower):
            data["nysc_status"] = NyscStatus.EXEMPTED
        elif re.search(r"ongoing|currently serving|corps member \(current|in view", lower):
            data["nysc_status"] = NyscStatus.ONGOING
        else:
            data["nysc_status"] = NyscStatus.COMPLETED
    else:
        data["nysc_status"] = ""

    explicit = [int(n) for n in re.findall(
        r"(\d{1,2})\+?\s*(?:years|yrs)\.?\s*(?:of\s+)?(?:[a-z-]+\s+){0,4}?experience", lower)]
    experience_text = sections.get("experience") or "\n".join(
        v for k, v in sections.items() if k not in {"education", "references"})
    computed, work_history = _experience_from_ranges(experience_text)
    data["years_experience"] = float(max(explicit)) if explicit and max(explicit) <= 45 else computed
    data["work_history"] = work_history[:10]
    if work_history:
        data["current_job_title"] = work_history[0]["job_title"]
        data["current_employer"] = work_history[0]["company"]

    data["tools"] = detect_terms(text, "tool", catalogue)
    detected_skills = detect_terms(text, "skill", catalogue)
    listed_skills = _list_items(sections.get("skills", ""), max_words=5)
    tool_norms = {normalize(t) for t in data["tools"]}
    data["skills"] = merge_unique(detected_skills, [s for s in listed_skills if normalize(s) not in tool_norms])[:40]
    data["certifications"] = merge_unique(
        detect_terms(text, "certification", catalogue), _list_items(sections.get("certifications", ""), max_words=12)
    )[:25]

    summary = sections.get("summary", "").strip()
    data["summary"] = re.sub(r"\s+", " ", summary)[:800]
    if data["highest_qualification"] or data["course"]:
        data["education_history"] = [{
            "qualification": EducationLevel(data["highest_qualification"]).label if data["highest_qualification"] else "",
            "course": data["course"], "institution": data["institution"],
            "year": str(data["graduation_year"] or ""),
        }]
    return data


def _parse_dob(value: str):
    value = value.strip().rstrip(".,")
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %B %Y", "%d %b %Y", "%B %d, %Y", "%b %d, %Y",
                "%d %B, %Y", "%Y-%m-%d", "%dth %B %Y", "%dst %B %Y", "%dnd %B %Y", "%drd %B %Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    m = re.match(r"(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]+),?\s+(\d{4})", value)
    if m:
        try:
            return datetime.strptime(f"{m.group(1)} {m.group(2)[:3]} {m.group(3)}", "%d %b %Y").date()
        except ValueError:
            return None
    return None


# ---------------------------------------------------------------------------
# AI extraction
# ---------------------------------------------------------------------------
CV_SCHEMA = {
    "type": "object",
    "properties": {
        "first_name": {"type": "string"},
        "last_name": {"type": "string"},
        "email": {"type": "string"},
        "phone": {"type": "string"},
        "gender": {"type": "string", "enum": ["", "male", "female"]},
        "date_of_birth": {"type": "string", "description": "YYYY-MM-DD or empty"},
        "location": {"type": "string"},
        "linkedin_url": {"type": "string"},
        "highest_qualification": {"type": "string", "enum": [""] + [c.value for c in EducationLevel]},
        "course": {"type": "string"},
        "institution": {"type": "string"},
        "graduation_year": {"type": "integer", "description": "0 if unknown"},
        "degree_class": {"type": "string", "enum": [""] + [c.value for c in DegreeClass]},
        "nysc_status": {"type": "string", "enum": [""] + [c.value for c in NyscStatus]},
        "years_experience": {"type": "number"},
        "current_employer": {"type": "string"},
        "current_job_title": {"type": "string"},
        "skills": {"type": "array", "items": {"type": "string"}},
        "tools": {"type": "array", "items": {"type": "string"}},
        "certifications": {"type": "array", "items": {"type": "string"}},
        "education_history": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "qualification": {"type": "string"},
                    "course": {"type": "string"},
                    "institution": {"type": "string"},
                    "year": {"type": "string"},
                },
                "required": ["qualification", "course", "institution", "year"],
                "additionalProperties": False,
            },
        },
        "work_history": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "job_title": {"type": "string"},
                    "company": {"type": "string"},
                    "start": {"type": "string"},
                    "end": {"type": "string"},
                    "description": {"type": "string"},
                },
                "required": ["job_title", "company", "start", "end", "description"],
                "additionalProperties": False,
            },
        },
        "summary": {"type": "string"},
    },
    "required": [
        "first_name", "last_name", "email", "phone", "gender", "date_of_birth", "location", "linkedin_url",
        "highest_qualification", "course", "institution", "graduation_year", "degree_class", "nysc_status",
        "years_experience", "current_employer", "current_job_title", "skills", "tools", "certifications",
        "education_history", "work_history", "summary",
    ],
    "additionalProperties": False,
}

CV_SYSTEM = (
    "You extract structured data from CVs/résumés submitted to a Nigerian manufacturing company. The document "
    "may be a scanned hard copy or a photo; read it carefully. Report only what the CV states — use an empty "
    "string, 0 or an empty list when something is not present. Qualification codes: ssce (WAEC/NECO), ond (OND/ND), "
    "nce, hnd, bsc (any bachelor's incl. BEng/BTech/LLB/MBBS), pgd, msc (any master's incl. MBA/MEng), phd. "
    "years_experience is total post-qualification work experience in years (count NYSC, exclude school "
    "attachments/IT unless no other work). Put software, systems, instruments and equipment in tools; "
    "competencies in skills; licences, professional memberships and certificates in certifications. "
    "Keep work_history descriptions under 40 words, most recent job first. summary: 2-3 sentence neutral "
    "profile of the candidate written in the third person."
)


def ai_extract(data: bytes, filename: str, extracted_text: str) -> dict | None:
    ext = os.path.splitext(filename or "")[1].lower()
    instruction = {"type": "text", "text": "Extract the candidate's details from this CV."}
    if ext == ".pdf":
        content = [ai.pdf_block(data), instruction]
    elif ext in CLAUDE_IMAGE_TYPES:
        content = [ai.image_block(data, CLAUDE_IMAGE_TYPES[ext]), instruction]
    elif ext in IMAGE_EXTENSIONS:  # TIFF/BMP scans: convert to PNG first
        from PIL import Image

        buf = io.BytesIO()
        Image.open(io.BytesIO(data)).convert("RGB").save(buf, format="PNG")
        content = [ai.image_block(buf.getvalue(), "image/png"), instruction]
    elif extracted_text.strip():
        content = [{"type": "text", "text": f"CV text:\n\n{extracted_text}"}, instruction]
    else:
        return None
    return ai.structured_request(system=CV_SYSTEM, content=content, schema=CV_SCHEMA, max_tokens=8000)


def _normalize_ai(result: dict) -> dict:
    result = dict(result)
    dob = result.get("date_of_birth") or ""
    try:
        result["date_of_birth"] = date.fromisoformat(dob) if dob else None
    except ValueError:
        result["date_of_birth"] = None
    result["graduation_year"] = result.get("graduation_year") or None
    for key in ("skills", "tools", "certifications"):
        result[key] = clean_list(result.get(key) or [])
    result["email"] = (result.get("email") or "").lower()
    return result


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------
LIST_FIELDS = ("skills", "tools", "certifications")


def parse_cv(data: bytes, filename: str) -> dict:
    """Return a dict of candidate fields plus ``text``, ``method`` and ``notes``."""
    extracted = extract_text(data, filename)
    notes = list(extracted.notes)
    heuristic = heuristic_extract(extracted.text)

    ai_data = None
    if ai.ai_enabled():
        ai_data = ai_extract(data, filename, extracted.text)
        if ai_data is None:
            notes.append("AI extraction unavailable; used rule-based extraction.")

    if ai_data:
        parsed = _normalize_ai(ai_data)
        for key, value in heuristic.items():
            if key in LIST_FIELDS:
                parsed[key] = merge_unique(list(parsed.get(key) or []), list(value or []))
            elif not parsed.get(key) and value:
                parsed[key] = value
        method = "ai"
    else:
        parsed = heuristic
        method = "heuristic" if extracted.text.strip() else ""

    text = extracted.text
    if not text.strip() and ai_data:
        # Scanned CV read by AI: keep a searchable text version for keyword matching.
        parts = [parsed.get("summary", ""), ", ".join(parsed.get("skills", [])), ", ".join(parsed.get("tools", [])),
                 ", ".join(parsed.get("certifications", []))]
        parts += [f"{w.get('job_title')} {w.get('company')} {w.get('description')}" for w in parsed.get("work_history", [])]
        parts += [f"{e.get('qualification')} {e.get('course')} {e.get('institution')}" for e in parsed.get("education_history", [])]
        text = "\n".join(p for p in parts if p)
    if not text.strip() and not ai_data and extracted.scanned:
        notes.append("Enable AI (ANTHROPIC_API_KEY) or install Tesseract OCR to read scanned CVs automatically.")

    parsed["text"] = text
    parsed["method"] = method
    parsed["notes"] = notes
    parsed["scanned"] = extracted.scanned
    return parsed
