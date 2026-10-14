"""Tests for loading the recruitment workbooks. All names and contacts here are made up."""

import os
import shutil
import tempfile
from datetime import date, datetime

from django.test import TestCase
from openpyxl import Workbook, load_workbook

from accounts.models import User
from candidates.models import Candidate, CBTResult
from core.models import Department, Employee
from core.workbook_import import WorkbookError, WorkbookImporter, degree_class, phones_in, qualification, to_date
from pipeline.models import Application, InterviewRecord, Onboarding
from requisitions.models import Requisition

CV_HEADERS = ["SN", "Name", "Phone Number", "Email", "Date of Birth", "Age", "Gender", "STATE", "LGA/Community Tag",
              "Relevant Experience", "Experience in Years", "Degree", "Discipline", "Year of Graduation", "Source",
              "Name of Institution", "Marks_Obtained", "CBT Status", "Interview Status"]
REQ_HEADERS = ["Req_Ref", "Requisition Date", "Name of Requisitor", "Company", "Department of Requisitor",
               "Functional Band", "Trainee/Experienced Band_Categories", "Recruitment Reason", "Location",
               "Role Required", "Employee Classsification", "Manpower Required", "Participants",
               "Shortlisted for Interview", "Rejected", "On-Hold", "Selected", "Onboarded  #", "In Progres",
               "Total Vacancies", "Status", "Interview Date", "Offer Date", "Offer Acceptance Date", "Closing Date",
               "Rec. Coordinator"]
IV_HEADERS = ["Ref_ID", "Remark", "Classification", "Name", "Gender", "Phone No.", "Email", "Date", "Panel Category",
              "Panellists", "Panel Member", "Community Query", "Panel Member2", "Score", "Remarks_Query",
              "Join Date", "Onboarding_Status"]
CBT_HEADERS = ["Surname", "Firstname", "Email", "Phone'# Check", "Date", "Category", "Total Score", "Exam Score",
               "Company", "Month-Year"]
STAFF_HEADERS = ["Date", "Emp No", "Emp Name", "Company Name", "Location ", "Join Date", "Emp Confirmation Date",
                 "Department", "Category", "Grade", "Designation", "Union Category", "Work Status", "Release Date",
                 "Reason of Release", "Consolidated Salary", "Birth Date", "Basic Qualification", "Specialization ",
                 "Bus. Unit", "Age", "Years to Retirement"]

PLACEHOLDER = "2348000000999"


def cv_rows():
    return [
        [1, "Ada Testperson", "08031234567, 09031234567", "ada.test@example.com", datetime(1995, 3, 13), 31,
         "Female", "Rivers State", "Eleme / Akpajo", "Yes", 2, "B.Sc", "Chemical Engineering", 2019, "Interna Portal",
         "Sample University", "Second class upper division", "Participated", "Interviewed"],
        [2, "Bayo Sample", "2348041112222", "bayo.sample@example.com", "22nd August 1995", "Not displayed",
         "Male", "Lagos", "Ikeja", "No", "8+", "HND", "Mechanical Engineering", "2093", "IFZ website",
         "Sample Polytechnic", "Upper Credit", "Not Participated", "Not Interviewed"],
        # Same email as Ada but a completely different name: kept apart, reported.
        [3, "Chidi Different", "08051112222", "ada.test@example.com", "Not Provided", None, "Male", "Imo", "",
         "Yes", 1, "OND", "#REF!", 2021, "PAC", "Sample College", "Lower Credit", "Not Participated",
         "Not Interviewed"],
    ]


def req_rows():
    return [
        ["Req_900", datetime(2025, 2, 1), "Jane Requester", "IFZ", "QA/QC", "Operations Support", "Lab_DET",
         "New Position", "Port Harcourt", "DET", "Trainee", 2, 4, 4, 2, 0, 2, 2, 0, 0, "Closed",
         datetime(2025, 2, 20), None, None, datetime(2025, 3, 3), "Coordinator One"],
        ["Req_901", datetime(2026, 8, 10), "John Requester", "IPL", "QAQC", "Engineering", "Experienced_Role",
         "Resignation / Replacement", "Port Harcourt", "Engineer", "Senior Staff", 1, 1, 1, 0, 0, 1, 0, 0, 1,
         "Pending", datetime(2026, 8, 25), None, None, None, "Coordinator Two"],
        [None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None,
         None, None, None, None, None, None, None, None],  # blank trailing row
    ]


def iv_rows():
    return [
        # Ada: hired on Req_900 (matched to her CV row by phone 0803..., email differs in case)
        ["Req_900", None, "DET", "Testperson Ada", "Female", "2348031234567", "ADA.TEST@example.com",
         datetime(2025, 2, 20), "Lab", "Lab Panel", "Panel A", None, None, 0.74, "Selected", datetime(2025, 5, 5),
         "Onboarded_IFL"],
        # Bayo: selected on Req_900 but not engaged (requisition closed)
        ["Req_900", None, "DET", "Bayo Sample", "Male", "2348041112222", "bayo.sample@example.com",
         datetime(2025, 2, 21), "Lab", "Lab Panel", "Panel A", None, None, 0.66, "Selected", None, "Not Engaged"],
        # Dan: selected on the pending Req_901 -> still active, waiting for a decision
        ["Req_901", "Interview_Shortlist", "Experienced", "Dan Interviewee", "Male", "2348062223333",
         "dan.iv@example.com", datetime(2026, 8, 25), "QAQC", "QA Panel", None, "Community", None, 0.81,
         "Selected", None, "Not Engaged"],
        # Interview for a requisition that is not in the tracker: kept on the profile only
        ["Req_12", None, "GET", "Dan Interviewee", "Male", "2348062223333", "dan.iv@example.com",
         datetime(2021, 5, 3), "Process", None, None, None, None, 0.3, "Rejected", None, "Not Engaged"],
    ]


def cbt_rows():
    return [
        ["Testperson", "Ada", "ada.test@example.com", "2348031234567", "2024-02-26 00:00:00", "Lab_GET", 150, 104,
         "IFL", "2024-02-01 00:00:00"],
        # exact repeat: loaded once
        ["Testperson", "Ada", "ada.test@example.com", "2348031234567", "2024-02-26 00:00:00", "Lab_GET", 150, 104,
         "IFL", "2024-02-01 00:00:00"],
        ["Only", "Cbt", "cbt.only@example.com", PLACEHOLDER, "2023-12-09 00:00:00", "Process_GET", 150, None, "IFL",
         "2023-12-01 00:00:00"],
        # date typed wrong: kept with no date, reported
        ["Wrongdate", "Eze", "eze.wrong@example.com", PLACEHOLDER, "0169-06-05 00:00:00", "MGT", None, None, "IFL",
         None],
        ["Fourth", "Person", "fourth@example.com", PLACEHOLDER, "2023-12-09 00:00:00", "Process_GET", 150, 70, "IFL",
         None],
        ["Fifth", "Human", "fifth@example.com", PLACEHOLDER, "2023-12-09 00:00:00", "Process_GET", 150, 71, "IFL",
         None],
        # Only contact is Ada's email, under another name: a separate person, found again on every run
        ["Other", "Zed", "ada.test@example.com", None, "2022-09-09 00:00:00", "Lab_DET", 150, 60, "IFL", None],
        # No contact at all, a repeated two-word name
        ["Olu", "Olu", None, None, "2021-05-26 00:00:00", "Mechanical_DET", 150, 55, "IFL", None],
        ["Olu", "Olu", None, None, "2022-05-26 00:00:00", "Mechanical_DET", 150, 58, "IFL", None],
    ]


def staff_rows():
    return [
        [datetime(2026, 8, 1), "SR9001", "Retiring Soonperson", "IEPL", "Port Harcourt", datetime(2006, 9, 13),
         datetime(2007, 3, 13), "QA/QC", "Staff Regular", "GRADE II D", "Technician", "Junior Staff", "Confirmed",
         None, None, 999999, datetime(1966, 12, 15), "ND", "Mechanical", "IPL", 59, 1],
        [datetime(2026, 8, 1), "FDT9002", "Young Trainee", "IEFCL", "Port Harcourt", datetime(2025, 2, 12),
         datetime(2026, 8, 12), "Human Resource", "Trainee-DET", "GRADE I/TRAINEES DET", "Trainee", "Trainee",
         "Trainee", None, None, 150000, datetime(2004, 11, 25), "ND", "Admin", "IFZ", 21, 39],
    ]


def write(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append(row)
    wb.save(path)


class WorkbookImportTests(TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.cv = os.path.join(self.dir, "CV_file.xlsx")
        self.master = os.path.join(self.dir, "Master.xlsx")
        write(self.cv, [("CV Dump", CV_HEADERS, cv_rows())])
        write(self.master, [
            ("Report Weekly Summary", ["Recruitment Report"], [["calculated"]]),
            ("Manpower_Requisition_Summary", REQ_HEADERS, req_rows()),
            ("Recruitment_Interview Records", IV_HEADERS, iv_rows()),
            ("Recruitment_CBTs", CBT_HEADERS, cbt_rows()),
            ("Retirement Forecast_3 Years Pro", STAFF_HEADERS, staff_rows()),
        ])
        self.reports = os.path.join(self.dir, "reports")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def run_import(self, replace=True, files=None):
        return WorkbookImporter(files or [self.master, self.cv], replace=replace, report_dir=self.reports,
                                log=lambda *a: None).run()

    def test_replace_removes_demo_data_and_keeps_real_accounts(self):
        dept = Department.objects.create(name="Demo Dept", code="DEMO")
        User.objects.create_user("hr.admin", "hr.admin@example.com", "x" * 10)
        owner = User.objects.create_user("owner", "owner@iefcl.example", "x" * 10, role="hr_admin")
        Candidate.objects.create(first_name="Demo", last_name="Person", email="demo@example.com")
        Requisition.objects.create(department=dept, title="Demo role")
        Employee.objects.create(staff_id="DEMO1", first_name="Demo", last_name="Staff")
        result = self.run_import()
        self.assertFalse(User.objects.filter(username="hr.admin").exists())
        self.assertTrue(User.objects.filter(pk=owner.pk).exists())
        self.assertFalse(Department.objects.filter(code="DEMO").exists())
        self.assertFalse(Candidate.objects.filter(email="demo@example.com").exists())
        self.assertFalse(Employee.objects.filter(staff_id="DEMO1").exists())
        self.assertFalse(Requisition.objects.filter(title="Demo role").exists())
        self.assertEqual(result.removed["demo sign-in accounts"], 1)

    def test_records_land_in_the_right_places(self):
        self.run_import()
        self.assertEqual(Requisition.objects.count(), 2)
        self.assertEqual(Employee.objects.count(), 2)
        # QA/QC and QAQC are one department; Human Resource is its own
        self.assertEqual(sorted(Department.objects.values_list("name", flat=True)), ["Human Resource", "QA/QC"])
        filled = Requisition.objects.get(reference="Req_900")
        self.assertEqual(filled.status, Requisition.Status.FILLED)
        self.assertEqual((filled.business_unit, filled.staff_category, filled.classification), ("IFZ", "Lab_DET", "Trainee"))
        self.assertEqual(filled.title, "Lab DET")
        self.assertEqual(filled.requested_by_name, "Jane Requester")
        self.assertEqual(filled.coordinator_name, "Coordinator One")
        self.assertEqual(filled.education_level, "")
        self.assertEqual(filled.created_at.date(), date(2025, 2, 1))
        pending = Requisition.objects.get(reference="Req_901")
        self.assertEqual(pending.status, Requisition.Status.APPROVED)  # active, not published
        self.assertIn(("Status", "Pending"), pending.tracker_facts)

        staff = Employee.objects.get(staff_id="SR9001")
        self.assertEqual((staff.company, staff.business_unit, staff.job_title), ("IEPL", "IPL", "Technician"))
        self.assertEqual(staff.retirement_date, date(2026, 12, 15))
        self.assertFalse(any("salary" in k.lower() for k in staff.source_record["values"]))

    def test_people_are_merged_without_duplicates(self):
        result = self.run_import()
        ada = Candidate.objects.get(phone="08031234567")
        self.assertEqual(ada.full_name, "Ada Testperson")
        self.assertEqual(ada.email, "ada.test@example.com")
        self.assertIn("09031234567", ada.other_contacts)
        self.assertEqual(ada.source_detail, "Internal Portal")  # "Interna Portal" typo corrected
        self.assertEqual((ada.state_of_origin, ada.lga), ("Rivers", "Eleme / Akpajo"))
        self.assertEqual((ada.highest_qualification, ada.degree_class), ("bsc", "2_1"))
        self.assertEqual(ada.cbt_results.count(), 1)  # the repeated CBT row is loaded once
        self.assertEqual(ada.interview_records.count(), 1)
        self.assertEqual(len(ada.source_records), 1)
        # Chidi shares Ada's email but not her name: a separate person, reported
        chidi = Candidate.objects.get(phone="08051112222")
        self.assertNotEqual(chidi.pk, ada.pk)
        self.assertEqual(chidi.course, "")  # "#REF!" is not a course
        self.assertTrue(result.shared_contacts)
        # Bayo: date written in words read; impossible graduation year left blank and reported
        bayo = Candidate.objects.get(email="bayo.sample@example.com")
        self.assertEqual(bayo.date_of_birth, date(1995, 8, 22))
        self.assertIsNone(bayo.graduation_year)
        self.assertEqual(float(bayo.years_experience), 8.0)
        self.assertTrue(any(i[3] == "Year of Graduation" for i in result.issues))
        # Dan appears only in interview records, twice
        dan = Candidate.objects.get(email="dan.iv@example.com")
        self.assertEqual(dan.source_detail, "Interview records")
        self.assertEqual(dan.community_status, "community")
        self.assertEqual(dan.interview_records.count(), 2)
        # One email per person and one phone per person
        self.assertEqual(Candidate.objects.filter(email="dan.iv@example.com").count(), 1)
        self.assertEqual(Candidate.objects.count(), 10)
        self.assertEqual(Candidate.objects.get(first_name="Olu").cbt_results.count(), 2)

    def test_placeholder_phone_is_never_used(self):
        result = self.run_import()
        self.assertIn("08000000999", result.placeholders)
        self.assertFalse(Candidate.objects.filter(phone="08000000999").exists())
        self.assertFalse(Candidate.objects.filter(other_contacts__contains="08000000999").exists())
        self.assertEqual(Candidate.objects.filter(source_detail="CBT records").count(), 6)
        wrong = Candidate.objects.get(email="eze.wrong@example.com").cbt_results.get()
        self.assertIsNone(wrong.test_date)

    def test_interview_outcomes_become_applications(self):
        self.run_import()
        ada = Application.objects.get(candidate__phone="08031234567")
        self.assertEqual((ada.stage, ada.status), ("hired", "hired"))
        self.assertEqual(ada.hired_at.date(), date(2025, 5, 5))
        bayo = Application.objects.get(candidate__email="bayo.sample@example.com")
        self.assertEqual((bayo.status, bayo.decision), ("not_engaged", "select"))
        dan = Application.objects.get(candidate__email="dan.iv@example.com")
        self.assertEqual((dan.stage, dan.status, dan.decision), ("decision", "active", "select"))
        self.assertEqual(Application.objects.count(), 3)
        self.assertFalse(Application.objects.filter(requisition__reference="Req_12").exists())
        record = InterviewRecord.objects.get(candidate__email="bayo.sample@example.com")
        self.assertEqual(float(record.score), 66.0)
        self.assertEqual(record.application, bayo)
        self.assertEqual(Onboarding.objects.count(), 0)

    def test_running_again_creates_no_duplicates(self):
        self.run_import()
        counts = [m.objects.count() for m in (Candidate, CBTResult, InterviewRecord, Application, Requisition, Employee,
                                               Department)]
        self.run_import(replace=True)
        self.assertEqual(counts, [m.objects.count() for m in (Candidate, CBTResult, InterviewRecord, Application,
                                                              Requisition, Employee, Department)])
        self.run_import(replace=False, files=[self.cv, self.master])  # merge mode, files in the other order
        self.assertEqual(counts, [m.objects.count() for m in (Candidate, CBTResult, InterviewRecord, Application,
                                                              Requisition, Employee, Department)])

    def test_report_and_wrong_files(self):
        result = self.run_import()
        wb = load_workbook(result.report_path)
        self.assertEqual(wb.sheetnames, ["Summary", "Values not used", "Shared contacts",
                                         "Same name, different contacts", "Duplicates merged"])
        with self.assertRaises(WorkbookError):
            WorkbookImporter([self.cv], replace=True, log=lambda *a: None).run()

    def test_pages_show_imported_records(self):
        self.run_import()
        hr = User.objects.create_user("owner", "owner@iefcl.example", "x" * 10, role="hr_admin")
        self.client.force_login(hr)
        ada = Candidate.objects.get(phone="08031234567")
        page = self.client.get(ada.get_absolute_url()).content.decode()
        for text in ("CBTs &amp; interviews", "Lab_GET", "As recorded in the CV list", "Eleme / Akpajo"):
            self.assertIn(text, page)
        req = Requisition.objects.get(reference="Req_900")
        page = self.client.get(req.get_absolute_url()).content.decode()
        for text in ("From the recruitment tracker", "Interview records", "Coordinator One", "Not scored"):
            self.assertIn(text, page)
        for url in ("/candidates/?record=cbt", "/candidates/?source=PAC", "/requisitions/?status=done&unit=IFZ",
                    "/staff/?view=all&unit=IPL", "/", f"/pipeline/applications/{ada.applications.get().pk}/"):
            self.assertEqual(self.client.get(url).status_code, 200, url)


class ParsingTests(TestCase):
    def test_dates(self):
        self.assertEqual(to_date("22nd August 1995"), date(1995, 8, 22))
        self.assertEqual(to_date("November 29th, 1991"), date(1991, 11, 29))
        self.assertEqual(to_date("28-Oct-1997"), date(1997, 10, 28))
        self.assertIsNone(to_date("32 years"))

    def test_qualifications_and_classes(self):
        self.assertEqual(qualification("B. Tech."), "bsc")
        self.assertEqual(qualification("ND"), "ond")
        self.assertEqual(qualification("Master of Business Administration"), "msc")
        self.assertEqual(qualification("Others"), "")
        self.assertEqual(degree_class("HND – Upper Credit; ND – Distinction"), "2_1")
        self.assertEqual(degree_class("Higher National Diploma"), "")

    def test_phones(self):
        self.assertEqual(phones_in("8160213803, 09044000000"), ["08160213803", "09044000000"])
        self.assertEqual(phones_in("+2348037954529"), ["08037954529"])
        self.assertEqual(phones_in("00000000000"), [])
