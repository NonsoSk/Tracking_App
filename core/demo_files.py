"""Generate simple CV files (PDF and Word) for demo data and tests."""

import io


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def simple_pdf(lines: list[str]) -> bytes:
    """A minimal, valid single-font PDF with one line of text per entry (multi-page)."""
    per_page = 48
    pages = [lines[i:i + per_page] for i in range(0, max(len(lines), 1), per_page)] or [[]]
    objects: list[bytes] = []
    # 1: catalog, 2: pages, 3: font, then page/content pairs
    kids = []
    page_objs = []
    for index, page_lines in enumerate(pages):
        page_id = 4 + index * 2
        content_id = page_id + 1
        kids.append(f"{page_id} 0 R")
        stream_lines = ["BT", "/F1 11 Tf", "14 TL", "56 790 Td"]
        for line in page_lines:
            stream_lines.append(f"({_pdf_escape(line.encode('latin-1', 'replace').decode('latin-1'))}) Tj T*")
        stream_lines.append("ET")
        stream = "\n".join(stream_lines).encode("latin-1", "replace")
        page_objs.append((page_id, f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                                   f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()))
        page_objs.append((content_id, b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.extend(body for _, body in page_objs)

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def simple_docx(lines: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for i, line in enumerate(lines):
        if i == 0:
            document.add_heading(line, level=1)
        elif line.isupper() and line.strip():
            document.add_heading(line.title(), level=2)
        else:
            document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def cv_lines(p: dict) -> list[str]:
    """Build realistic CV text from a small profile dict."""
    lines = [f"{p['first']} {p['last']}".upper(), p.get("address", "Port Harcourt, Rivers State"),
             f"Email: {p['email']} | Phone: {p['phone']}", "", "PROFESSIONAL SUMMARY", p["summary"], "",
             "WORK EXPERIENCE"]
    for job in p.get("jobs", []):
        lines += [f"{job['title']}, {job['company']}", f"{job['start']} - {job['end']}"]
        lines += [f"- {b}" for b in job.get("bullets", [])]
    if p.get("nysc"):
        lines += [f"NYSC Corps Member, {p['nysc']}"]
    lines += ["", "EDUCATION", f"{p['degree']} {p['course']}, {p['school']}   {p['year']}"]
    if p.get("class"):
        lines.append(p["class"])
    lines += ["WAEC (SSCE) 2008", "", "SKILLS", ", ".join(p.get("skills", [])), "", "CERTIFICATIONS"]
    lines += p.get("certs", []) or ["-"]
    lines += ["", "REFERENCES", "Available on request"]
    return lines
