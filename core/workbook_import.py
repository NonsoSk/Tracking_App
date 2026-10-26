"""Load IEFCL's recruitment workbooks into the portal.

Two workbooks are read:

* the CV file ("CV Dump" sheet): the pool of applicants;
* the recruitment master workbook: the manpower requisition tracker, the interview
  records, the CBT results and the staff list (retirement forecast).

Sheets are found by their column headings, so the files can be given in any order and
renamed. The calculated sheets (weekly summary, HR metrics, analysis, dashboard) hold no
records of their own and are not read.

With ``replace=True`` every record already in the portal (including the demo data) is
removed first, so the portal holds only what is in the two workbooks. Settings are kept:
the selection-report criteria, the document checklists and skill vocabulary.

Without it the run is a refresh (see core/excel_link.py): new rows are added, rows changed in
Excel update their records, and what the hiring team did in the portal is kept. A requisition's
status follows the tracker only when the tracker's Status cell changed since the last run, and an
application follows the interview records only while nobody has worked on it in the portal.
A refresh may be given just one of the workbooks.

People who appear in several sheets (same email or phone number, and names that share a
word) become one candidate. A phone number or email used by more than three different
people is treated as a placeholder and never used to match anyone. Salaries are never read.
"""

from __future__ import annotations

import os
import re
from collections import Counter, defaultdict
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal

from django.db import transaction
from django.utils import timezone

from accounts.models import User
from candidates.models import Candidate, CandidateDocument, CBTResult, CommunityStatus
from core.choices import DegreeClass, EducationLevel, EmploymentType
from core.models import Department, Employee, Notification
from pipeline.models import (
    Activity, Application, InterviewRecord, MedicalCheck, Offer, Onboarding, OnboardingTask, Stage,
)
from pipeline.services import DEFAULT_ONBOARDING_TASKS
from requisitions.models import JobRole, Requisition, RequisitionNote


class WorkbookError(Exception):
    """The files are not the expected workbooks."""


# Headings that identify each kind of sheet (compared in lower case, spaces collapsed).
SHEET_SIGNATURES = {
    "cv": {"name", "phone number", "email", "degree", "discipline", "source"},
    "requisitions": {"req_ref", "manpower required", "department of requisitor", "status"},
    "interviews": {"ref_id", "name", "panel category", "remarks_query"},
    "cbt": {"surname", "firstname", "exam score", "category"},
    "staff": {"emp no", "emp name", "birth date", "department"},
}
SHEET_LABELS = {
    "cv": "CV list", "requisitions": "Manpower requisition tracker", "interviews": "Interview records",
    "cbt": "CBT results", "staff": "Staff list (retirement forecast)",
}

# Values that mean "nothing recorded".
BLANKS = {"", "nan", "none", "null", "n/a", "na", "nil", "-", "--", "not displayed", "not displayrd",
          "not provided", "not indicated", "not available", "nill",
          # Excel error values left by broken formulas
          "#ref!", "#n/a", "#value!", "#div/0!", "#name?", "#null!", "#num!"}

PLACEHOLDER_LIMIT = 3  # a contact shared by more people than this is not a personal contact

# Usernames created by the demo seed (seed_demo).
DEMO_USERNAMES = {"admin", "hr.admin", "recruiter", "hod.production", "hod.maintenance", "hod.ict", "hod.hse",
                  "manager.lab", "interviewer", "management", "onboarding", "staff.referrer"}

# Department spellings that mean the same department.
DEPARTMENT_ALIASES = {"hr": "humanresources", "humanresource": "humanresources", "it": "informationtechnology"}

SOURCE_SPELLINGS = {"interna portal": "Internal Portal", "internal portal": "Internal Portal"}


# --------------------------------------------------------------------------- helpers
def header_key(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip().lower()


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ").replace("_x000d_", "")).strip()
    return "" if text.lower() in BLANKS else text


def nice_name(text: str) -> str:
    """Title-case names typed all in capitals or all in lower case; keep everything else."""
    text = clean(text)
    if text and (text.isupper() or text.islower()):
        return " ".join(part.capitalize() for part in text.split())
    return text


def split_name(full: str) -> tuple[str, str]:
    parts = nice_name(full).split()
    if not parts:
        return "", ""
    return parts[0][:80], " ".join(parts[1:])[:80]


def full_name_key(name: str) -> str:
    """Words of a full name in alphabetical order, for names of at least two words."""
    words = sorted(w for w in re.findall(r"[a-z]+", str(name or "").lower()) if w not in {"nan", "none"})
    return " ".join(words) if len(words) >= 2 else ""


def name_tokens(name: str) -> frozenset:
    return frozenset(t for t in re.findall(r"[a-z]{3,}", str(name or "").lower()) if t not in {"nan", "none"})


DATE_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%b-%Y", "%d-%B-%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y",
                "%d %b %Y", "%d %B %Y", "%b %d %Y", "%B %d %Y", "%m/%d/%Y", "%m-%d-%Y")


def to_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value)
    if not text:
        return None
    # "22nd August 1995", "November 29th, 1991"
    text = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", text, flags=re.I).replace(",", " ")
    text = re.sub(r"\s+", " ", text).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def to_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = clean(value).replace(",", "")
    match = re.match(r"^(\d+(?:\.\d+)?)\+?$", text)
    return Decimal(match.group(1)) if match else None


def emails_in(value) -> list[str]:
    text = clean(value).lower()
    return [e.strip(".") for e in re.split(r"[\s,;/]+", text) if re.match(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", e.strip("."))]


def phones_in(value) -> list[str]:
    """Nigerian numbers as 11 digits starting with 0 (234... and missing leading zeros fixed)."""
    text = clean(value)
    out = []
    for piece in re.split(r"[,/;]|\s{2,}| and ", text):
        digits = re.sub(r"\D", "", piece)
        if digits.startswith("234"):
            digits = digits[3:]
        digits = digits.lstrip("0")
        if len(digits) == 10 and len(set(digits)) > 2:
            out.append("0" + digits)
    return out


def jsonable(value):
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.replace("_x000d_", "").strip()
    return value


def aware(d):
    if not d:
        return None
    return timezone.make_aware(datetime.combine(d, time(9, 0)))


QUALIFICATION_RULES = [
    (EducationLevel.PHD, ("phd", "doctor")),
    (EducationLevel.MSC, ("msc", "m.sc", "mtech", "m.tech", "meng", "m.eng", "mba", "master", "m.ed")),
    (EducationLevel.PGD, ("pgd", "post graduate", "postgraduate")),
    (EducationLevel.HND, ("hnd", "higher national")),
    (EducationLevel.BSC, ("bsc", "b.sc", "beng", "b.eng", "btech", "b.tech", "b. tech", "b. eng", "b.ed", "bed", "ba",
                          "bachelor", "bse", "b.mls", "llb", "b.a", "b.pharm")),
    (EducationLevel.OND, ("ond", "nd", "national diploma")),
    (EducationLevel.NCE, ("nce",)),
    (EducationLevel.SSCE, ("ssce", "waec", "neco")),
]


def qualification(value) -> str:
    text = clean(value).lower()
    if not text or text in {"others", "other"}:
        return ""
    squashed = text.replace(" ", "")
    for level, words in QUALIFICATION_RULES:
        for word in words:
            w = word.strip()
            if text == w or squashed == w.replace(" ", "") or squashed.startswith(w.replace(" ", "")):
                return level
    return ""


def degree_class(value) -> str:
    text = clean(value).lower()
    if not text:
        return ""
    if text == "pass":
        return DegreeClass.PASS
    words = [("first", DegreeClass.FIRST), ("distinction", DegreeClass.FIRST), ("upper", DegreeClass.SECOND_UPPER),
             ("lower", DegreeClass.SECOND_LOWER), ("third", DegreeClass.THIRD)]
    found = [(text.find(word), cls) for word, cls in words if word in text]
    return min(found)[1] if found else ""


def gender(value) -> str:
    text = clean(value).lower()
    return {"male": "male", "m": "male", "female": "female", "f": "female"}.get(text, "")


def yes_no(value):
    text = clean(value).lower()
    if text in {"yes", "y"}:
        return True
    if text in {"no", "non", "n"}:
        return False
    return None


def state_name(value) -> str:
    text = nice_name(value)
    return re.sub(r"\s+state$", "", text, flags=re.I).strip()


INTERVIEW_RESULTS = {
    "selected": InterviewRecord.Result.SELECTED, "rejected": InterviewRecord.Result.REJECTED,
    "on-hold": InterviewRecord.Result.ON_HOLD, "on hold": InterviewRecord.Result.ON_HOLD,
    "rescheduled": InterviewRecord.Result.RESCHEDULED, "notfound": InterviewRecord.Result.NOT_FOUND,
    "not found": InterviewRecord.Result.NOT_FOUND, "declined": InterviewRecord.Result.DECLINED,
}

REASONS = [
    ("retirement", Requisition.Reason.RETIREMENT),
    ("resignation", Requisition.Reason.RESIGNATION),
    ("death", Requisition.Reason.RESIGNATION),
    ("medical", Requisition.Reason.RESIGNATION),
    ("expat", Requisition.Reason.TRANSFER),
    ("project", Requisition.Reason.TEMPORARY),
]


# --------------------------------------------------------------------------- data classes
@dataclass
class Sheet:
    kind: str
    file: str
    title: str
    labels: dict  # lower-case heading -> heading as written in the sheet
    rows: list  # (excel row number, dict of lower-case heading -> value)

    def original(self, row: dict) -> dict:
        """The row keyed by the sheet's own headings, ready to store as JSON."""
        return {self.labels.get(h, h): jsonable(v) for h, v in row.items()}


@dataclass
class PersonRow:
    kind: str  # cv | interview | cbt | existing
    ref: object  # (sheet, row) or candidate pk
    name: str
    emails: list
    phones: list
    data: dict = field(default_factory=dict)
    when: date | None = None


@dataclass
class ImportResult:
    counts: Counter = field(default_factory=Counter)
    removed: Counter = field(default_factory=Counter)
    issues: list = field(default_factory=list)  # (file, sheet, row, column, value, note)
    shared_contacts: list = field(default_factory=list)
    placeholders: dict = field(default_factory=dict)
    merged: list = field(default_factory=list)  # (name, rows, sources)
    same_names: list = field(default_factory=list)  # (name, contacts of each person)
    report_path: str = ""


def snapshot(instance) -> dict:
    """Field values, to tell whether a record really changed."""
    return {f.attname: deepcopy(getattr(instance, f.attname)) for f in instance._meta.concrete_fields}


# --------------------------------------------------------------------------- reading
def read_sheets(paths, *, require_all=True, skipped=None) -> dict[str, Sheet]:
    """The record sheets in these files, by kind. A second sheet of a kind already found is listed in ``skipped``."""
    from openpyxl import load_workbook

    found: dict[str, Sheet] = {}
    for path in paths:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any unreadable file
            raise WorkbookError(f"Could not open {os.path.basename(path)}: {exc}") from exc
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration:
                continue
            headers = [header_key(h) for h in raw_headers]
            labels = {header_key(h): re.sub(r"\s+", " ", str(h).replace("\xa0", " ")).strip()
                      for h in raw_headers if header_key(h)}
            kind = next((k for k, sig in SHEET_SIGNATURES.items() if sig <= set(headers)), None)
            if not kind:
                continue
            if kind in found:
                if skipped is not None:
                    skipped.append((os.path.basename(path), ws.title, kind, found[kind]))
                continue
            data = []
            for number, values in enumerate(rows, start=2):
                record = {h: v for h, v in zip(headers, values) if h}
                if any(clean(v) for v in record.values() if not isinstance(v, bool)):
                    data.append((number, record))
            found[kind] = Sheet(kind, os.path.basename(path), ws.title, labels, data)
        wb.close()
    missing = [SHEET_LABELS[k] for k in SHEET_SIGNATURES if k not in found]
    if missing and (require_all or not found):
        names = ", ".join(os.path.basename(str(p)) for p in paths)
        if not found:
            raise WorkbookError(f"No recruitment sheets were found in {names}. Is this the CV file or the "
                                "recruitment master workbook?")
        raise WorkbookError("These sheets were not found in the files given: " + ", ".join(missing) + ".")
    return found


# --------------------------------------------------------------------------- importer
class WorkbookImporter:
    def __init__(self, paths, *, replace=False, report_dir=None, log=print, require_all=True):
        self.paths = list(paths)
        self.replace = replace
        self.require_all = require_all or replace  # never empty the portal to load only part of the data
        self.report_dir = report_dir
        self.log = log
        self.result = ImportResult()
        self.today = timezone.localdate()

    # -- bookkeeping
    def issue(self, sheet: Sheet, row, column, value, note):
        self.result.issues.append((sheet.file, sheet.title, row, column, jsonable(value), note))

    def event_date(self, sheet, row, column, value):
        d = to_date(value)
        if value not in (None, "") and clean(value) and d is None:
            self.issue(sheet, row, column, value, "Not a date; left blank")
        if d and not (date(2000, 1, 1) <= d <= date(self.today.year + 1, 12, 31)):
            self.issue(sheet, row, column, value, "Date out of range; left blank")
            return None
        return d

    def birth_date(self, sheet, row, column, value):
        d = to_date(value)
        if clean(value) and d is None:
            self.issue(sheet, row, column, value, "Not a date; left blank")
        if d and not (date(1940, 1, 1) <= d <= date(self.today.year - 14, 12, 31)):
            self.issue(sheet, row, column, value, "Birth date out of range; left blank")
            return None
        return d

    # -- main
    def run(self) -> ImportResult:
        self.log("Reading the workbooks...")
        skipped = []
        sheets = read_sheets(self.paths, require_all=self.require_all, skipped=skipped)
        self.sheets = sheets
        for file, title, kind, used in skipped:
            self.result.issues.append((file, title, "", "", "", f"Not read: a second {SHEET_LABELS[kind]} sheet. Only "
                                       f"'{used.title}' ({used.file}) was read. Keep one sheet of each kind per "
                                       "workbook, or load this one separately."))
            self.result.counts["sheets not read (a second sheet of the same kind; see report)"] += 1
        for kind, sheet in sheets.items():
            self.log(f"  {SHEET_LABELS[kind]}: {len(sheet.rows):,} rows ({sheet.file} > {sheet.title})")
        with transaction.atomic():
            if self.replace:
                self.log("Removing every existing record (demo data included)...")
                self.purge()
            self.departments = {}
            self.requisitions = {}
            self.load_departments(sheets)
            if "staff" in sheets:
                self.load_staff(sheets["staff"])
            if "requisitions" in sheets:
                self.load_requisitions(sheets["requisitions"])
            if sheets.keys() & {"cv", "interviews", "cbt"}:
                self.load_people(sheets)
        if self.report_dir:
            self.result.report_path = write_report(self.result, self.report_dir)
        return self.result

    # -- purge
    def purge(self):
        removed = self.result.removed
        for doc in CandidateDocument.objects.exclude(file=""):
            doc.file.delete(save=False)
        for offer in Offer.objects.exclude(letter=""):
            offer.letter.delete(save=False)
        for medical in MedicalCheck.objects.exclude(report=""):
            medical.report.delete(save=False)
        for model in (Notification, Activity, InterviewRecord, CBTResult, Application, CandidateDocument, Candidate,
                      RequisitionNote, Requisition, JobRole, Employee):
            removed[str(model._meta.verbose_name_plural)] += model.objects.all().delete()[1].get(model._meta.label, 0)
        demo_users = User.objects.filter(username__in=DEMO_USERNAMES, email__iendswith="@example.com")
        removed["demo sign-in accounts"] += demo_users.count()
        demo_users.delete()
        removed["departments"] += Department.objects.all().delete()[1].get(Department._meta.label, 0)

    # -- departments
    def dept_key(self, name: str) -> str:
        key = re.sub(r"[^a-z0-9]", "", name.lower())
        return DEPARTMENT_ALIASES.get(key, key)

    def load_departments(self, sheets):
        spellings = defaultdict(Counter)
        for kind, column in (("staff", "department"), ("requisitions", "department of requisitor")):
            for _, row in (sheets[kind].rows if kind in sheets else []):
                name = clean(row.get(column))
                if name:
                    spellings[self.dept_key(name)][name] += 1
        existing = {self.dept_key(d.name): d for d in Department.objects.all()}
        codes = set(Department.objects.values_list("code", flat=True))
        for key, names in spellings.items():
            if key in existing:
                self.departments[key] = existing[key]
                continue
            name = names.most_common(1)[0][0][:150]
            code = re.sub(r"[^A-Z0-9]", "", name.upper())[:16] or "DEPT"
            base, n = code, 2
            while code in codes:
                code, n = f"{base[:16]}{n}", n + 1
            codes.add(code)
            self.departments[key] = Department.objects.create(name=name, code=code)
            self.result.counts["departments"] += 1
            if len(names) > 1:
                self.result.issues.append(("", "Departments", "", "Department", ", ".join(names),
                                           f"Spellings treated as one department: {name}"))

    def department(self, name):
        name = clean(name)
        return self.departments.get(self.dept_key(name)) if name else None

    # -- staff
    def load_staff(self, sheet: Sheet):
        self.log(f"Importing staff ({len(sheet.rows):,})...")
        existing = {e.staff_id: e for e in Employee.objects.all()}
        new, updated, seen = [], [], set()
        for number, row in sheet.rows:
            staff_id = clean(row.get("emp no"))
            if not staff_id and not clean(row.get("emp name")):
                continue
            if not staff_id:
                self.issue(sheet, number, "Emp No", "", "No staff number; row skipped")
                continue
            if staff_id in seen:
                self.issue(sheet, number, "Emp No", staff_id, "Staff number appears twice; second row skipped")
                continue
            seen.add(staff_id)
            first, last = split_name(row.get("emp name"))
            release = to_date(row.get("release date"))
            reason = clean(row.get("reason of release")).lower()
            status = Employee.Status.ACTIVE
            if release or reason:
                status = next((s for word, s in (("retire", Employee.Status.RETIRED), ("resign", Employee.Status.RESIGNED),
                                                  ("transfer", Employee.Status.TRANSFERRED), ("terminat", Employee.Status.TERMINATED),
                                                  ("dismiss", Employee.Status.TERMINATED)) if word in reason),
                              Employee.Status.RESIGNED)
            values = {
                "first_name": first, "last_name": last,
                "department": self.department(row.get("department")),
                "job_title": clean(row.get("designation"))[:150],
                "grade": clean(row.get("grade"))[:40],
                "date_of_birth": self.birth_date(sheet, number, "Birth Date", row.get("birth date")),
                "date_of_employment": to_date(row.get("join date")),
                "company": clean(row.get("company name"))[:60],
                "business_unit": clean(row.get("bus. unit"))[:20],
                "location": clean(row.get("location"))[:120],
                "category": clean(row.get("category"))[:60],
                "work_status": clean(row.get("work status"))[:40],
                "status": status, "exit_date": release,
                "source_record": {"file": sheet.file, "sheet": sheet.title, "row": number,
                                  "values": {k: v for k, v in sheet.original(row).items() if "salary" not in k.lower()}},
            }
            employee = existing.get(staff_id)
            if employee:
                if all(getattr(employee, key) == value for key, value in values.items()):
                    continue  # unchanged since the last run
                for key, value in values.items():
                    setattr(employee, key, value)
                employee.retirement_date = employee.compute_retirement_date()
                updated.append(employee)
            else:
                employee = Employee(staff_id=staff_id[:30], **values)
                employee.retirement_date = employee.compute_retirement_date()
                new.append(employee)
        Employee.objects.bulk_create(new, batch_size=500)
        if updated:
            Employee.objects.bulk_update(updated, [f for f in values if f] + ["retirement_date"], batch_size=500)
        self.result.counts["staff added"] += len(new)
        self.result.counts["staff updated"] += len(updated)

    # -- requisitions
    def load_requisitions(self, sheet: Sheet):
        self.log(f"Importing requisitions ({len(sheet.rows):,})...")
        self.requisitions = {}
        existing = {r.reference.lower(): r for r in Requisition.objects.exclude(reference="")}
        seen = set()
        for number, row in sheet.rows:
            ref = clean(row.get("req_ref"))
            dept = self.department(row.get("department of requisitor"))
            role = clean(row.get("role required"))
            if not (ref or dept or role):
                continue
            if not dept or not role:
                self.issue(sheet, number, "Department of Requisitor / Role Required", ref,
                           "No department or role; row skipped")
                continue
            if ref and ref.lower() in seen:
                self.issue(sheet, number, "Req_Ref", ref, "Reference appears twice; second row skipped")
                continue
            if not ref:
                self.issue(sheet, number, "Req_Ref", "", "No requisition reference; the portal gave it one")
            seen.add(ref.lower())
            category = clean(row.get("trainee/experienced band_categories"))
            classification = clean(row.get("employee classsification") or row.get("employee classification"))
            prefix = category.split("_")[0] if "_" in category else ""
            title = f"{prefix} {role}" if role.upper() in {"GET", "DET"} and prefix and prefix.lower() != "experienced" else role
            reason_text = clean(row.get("recruitment reason"))
            reason = next((r for word, r in REASONS if word in reason_text.lower()), Requisition.Reason.NEW_POSITION)
            if classification.lower() == "trainee":
                employment = EmploymentType.TRAINEE
            elif classification.lower() in {"contract", "retainer"}:
                employment = EmploymentType.CONTRACT
            else:
                employment = EmploymentType.EXPERIENCED
            positions = to_number(row.get("manpower required"))
            if positions is None:
                self.issue(sheet, number, "Manpower Required", row.get("manpower required"), "Not a number; set to 1")
                positions = 1
            onboarded = to_number(row.get("onboarded #")) or 0
            tracker_status = clean(row.get("status")).lower()
            if tracker_status == "closed":
                status = Requisition.Status.FILLED if onboarded >= positions else Requisition.Status.CLOSED
            else:
                status = Requisition.Status.APPROVED
                if tracker_status not in {"pending", "in progress"}:
                    self.issue(sheet, number, "Status", row.get("status"), "Unknown status; treated as open (approved)")
            req_date = self.event_date(sheet, number, "Requisition Date", row.get("requisition date"))
            closing = self.event_date(sheet, number, "Closing Date", row.get("closing date"))
            values = {
                "department": dept, "title": title[:150], "positions": int(positions), "reason": reason,
                "reason_detail": reason_text[:150], "employment_type": employment,
                "location": clean(row.get("location"))[:120], "education_level": "",
                "status": status, "business_unit": clean(row.get("company"))[:20],
                "functional_band": clean(row.get("functional band"))[:60], "staff_category": category[:60],
                "classification": classification[:40],
                "requested_by_name": nice_name(row.get("name of requisitor"))[:150],
                "coordinator_name": nice_name(row.get("rec. coordinator"))[:150],
                "submitted_at": aware(req_date), "closed_at": aware(closing) if status != Requisition.Status.APPROVED else None,
                "source_record": {"file": sheet.file, "sheet": sheet.title, "row": number,
                                  "key": ref or f"row:{req_date}|{clean(row.get('company'))}|{dept.name}|{role}",
                                  "values": sheet.original(row)},
            }
            key = values["source_record"]["key"]
            req = existing.get(ref.lower()) if ref else Requisition.objects.filter(source_record__key=key).first()
            if req and not self.replace:
                self.keep_portal_changes(req, values, tracker_status)
            if req:
                if any(getattr(req, k) != v for k, v in values.items()):
                    for k, v in values.items():
                        setattr(req, k, v)
                    req.save()
                    self.result.counts["requisitions updated"] += 1
            else:
                req = Requisition(reference=ref[:20], **values)
                req.save()
                self.result.counts["requisitions added"] += 1
            if req_date and req.created_at != aware(req_date):
                Requisition.objects.filter(pk=req.pk).update(created_at=aware(req_date))
            if ref:
                self.requisitions[ref.lower()] = req

    def keep_portal_changes(self, req, values, tracker_status):
        """On a refresh, keep what the hiring team set in the portal unless Excel changed the same thing."""
        if req.education_level:
            values.pop("education_level")
        before = next((clean(v).lower() for k, v in (req.source_record or {}).get("values", {}).items()
                       if header_key(k) == "status"), None)
        closed_now = values["status"] != Requisition.Status.APPROVED
        portal_closed = req.status in {Requisition.Status.FILLED, Requisition.Status.CLOSED}
        if tracker_status == before:
            follow = closed_now and portal_closed  # still closed: filled or not may have changed
        elif closed_now:
            follow = req.status != Requisition.Status.CANCELLED
        else:
            follow = portal_closed  # re-opened in the tracker
        if not follow:
            values.pop("status")
            values.pop("closed_at")
        elif values["status"] != req.status:
            self.result.counts["requisition statuses changed by the tracker"] += 1

    # -- people
    def load_people(self, sheets):
        for req in Requisition.objects.exclude(reference=""):
            self.requisitions.setdefault(req.reference.lower(), req)
        empty = Sheet("", "", "", {}, [])
        cv, interviews, cbt = (sheets.get(k, empty) for k in ("cv", "interviews", "cbt"))
        rows: list[PersonRow] = []
        def has_person(row, *columns):
            return any(clean(row.get(c)) for c in columns)

        for number, row in cv.rows:
            if not has_person(row, "name", "email", "phone number"):
                continue
            rows.append(PersonRow("cv", (cv, number), clean(row.get("name")), emails_in(row.get("email")),
                                  phones_in(row.get("phone number")), row))
        for number, row in interviews.rows:
            if not has_person(row, "name", "email", "phone no."):
                continue
            rows.append(PersonRow("interview", (interviews, number), clean(row.get("name")),
                                  emails_in(row.get("email")), phones_in(row.get("phone no.")), row,
                                  to_date(row.get("date"))))
        for number, row in cbt.rows:
            if not has_person(row, "firstname", "surname", "email"):
                continue
            name = f"{clean(row.get('firstname'))} {clean(row.get('surname'))}".strip()
            rows.append(PersonRow("cbt", (cbt, number), name, emails_in(row.get("email")),
                                  phones_in(row.get("phone'# check") or row.get("phone")), row, to_date(row.get("date"))))
        if not self.replace:
            for c in Candidate.objects.all().only("id", "first_name", "last_name", "email", "phone", "other_contacts"):
                rows.append(PersonRow("existing", c.pk, c.full_name, emails_in(f"{c.email} {c.other_contacts}"),
                                      phones_in(f"{c.phone}, {c.other_contacts}")))

        self.log(f"Matching people across the sheets ({len(rows):,} rows)...")
        groups = self.group_people(rows)
        self.log(f"  {len(groups):,} different people found")
        self.save_people(groups, sheets)

    def group_people(self, rows):
        # Contacts shared by several unrelated names are placeholders (an office number, say).
        users = defaultdict(list)
        for row in rows:
            for key in [("e", e) for e in row.emails] + [("p", p) for p in row.phones]:
                users[key].append(name_tokens(row.name))
        placeholders = set()
        for key, names in users.items():
            if len(names) <= PLACEHOLDER_LIMIT:
                continue
            clusters = []
            for n in names:
                for cl in clusters:
                    if cl & n:
                        cl |= n
                        break
                else:
                    clusters.append(set(n))
            if len(clusters) > PLACEHOLDER_LIMIT:
                placeholders.add(key)
                self.result.placeholders[key[1]] = len(names)
        self.placeholders = {k[1] for k in placeholders}

        parent = list(range(len(rows)))

        def find(i):
            while parent[i] != i:
                parent[i] = parent[parent[i]]
                i = parent[i]
            return i

        # A contact can be shared by people with different names (a family phone, a typing slip), so each row
        # joins the earlier holder of the contact whose name matches, never one with an unrelated name.
        holders = defaultdict(list)
        reported = set()
        for i, row in enumerate(rows):
            mine = name_tokens(row.name)
            for key in [("e", e) for e in row.emails] + [("p", p) for p in row.phones]:
                if key in placeholders:
                    continue
                match = next((j for j in holders[key]
                              if not mine or not name_tokens(rows[j].name) or mine & name_tokens(rows[j].name)), None)
                if match is not None:
                    parent[find(i)] = find(match)
                elif holders[key] and row.kind != "existing" and (key, mine) not in reported:
                    reported.add((key, mine))
                    self.result.shared_contacts.append((key[1], self.describe(rows[holders[key][0]]),
                                                        self.describe(row)))
                holders[key].append(i)
        groups = defaultdict(list)
        for i in range(len(rows)):
            groups[find(i)].append(rows[i])
        groups = list(groups.values())

        # Rows with no usable contact join a person only if exactly one person has the same full name.
        def usable(r):
            return [e for e in r.emails if e not in self.placeholders] or [p for p in r.phones if p not in self.placeholders]

        by_name = defaultdict(dict)
        merged_groups, loose = [], []
        for g in groups:
            if any(usable(r) for r in g):
                merged_groups.append(g)
                for r in g:
                    if full_name_key(r.name):
                        by_name[full_name_key(r.name)][id(g)] = g
            else:
                loose.append(g)
        alone = defaultdict(list)
        for g in loose:
            key = full_name_key(g[0].name)
            matches = by_name.get(key, {})
            if key and len(matches) == 1:
                next(iter(matches.values())).extend(g)
            elif key:
                alone[key].extend(g)  # same full name and no contact anywhere: one person
            else:
                merged_groups.append(g)
        merged_groups.extend(alone.values())
        return merged_groups

    def describe(self, row: PersonRow) -> str:
        if row.kind == "existing":
            return f"{row.name} (already in the portal)"
        sheet, number = row.ref
        return f"{row.name} ({sheet.title} row {number})"

    def save_people(self, groups, sheets):
        existing = {c.pk: c for c in Candidate.objects.filter(
            pk__in=[r.ref for g in groups for r in g if r.kind == "existing"])}
        order = {"existing": 0, "cv": 1, "interview": 2, "cbt": 3}
        new_people, plan = [], []
        for group in groups:
            group.sort(key=lambda r: (order[r.kind], -(r.when or date.min).toordinal()))
            current = next((existing[r.ref] for r in group if r.kind == "existing"), None)
            person = current or Candidate(source=Candidate.Source.IMPORT, consent_given=False, years_experience=None)
            before = snapshot(person)
            self.fill_person(person, group)
            if current is None:
                new_people.append(person)
            elif any(r.kind != "existing" for r in group) and snapshot(person) != before:
                person.save()
                self.result.counts["candidates updated"] += 1
            plan.append((person, group))
            imported = [r for r in group if r.kind != "existing"]
            if len(imported) > 1:
                kinds = Counter(r.kind for r in imported)
                self.result.merged.append((person.full_name, len(imported),
                                           ", ".join(f"{n} {k}" for k, n in kinds.items())))
        self.log(f"Saving {len(new_people):,} candidates...")
        Candidate.objects.bulk_create(new_people, batch_size=500)
        if new_people and new_people[0].pk is None:
            raise WorkbookError("The database did not return new record numbers; please update SQLite.")
        # Show when each person first appeared, not the day of the import.
        new_ids = {id(p) for p in new_people}
        dated = []
        for person, group in plan:
            first_seen = min((r.when for r in group if r.when), default=None)
            if first_seen and id(person) in new_ids:
                person.created_at = aware(first_seen)
                dated.append(person)
        Candidate.objects.bulk_update(dated, ["created_at"], batch_size=500)
        by_name = defaultdict(list)
        for person, _ in plan:
            by_name[" ".join(sorted(person.full_name.lower().split()))].append(person)
        for people in by_name.values():
            if len(people) > 1:
                self.result.same_names.append((people[0].full_name, " | ".join(
                    f"{p.email or '-'} {p.phone or '-'} ({p.source_detail or 'spreadsheet'})" for p in people)))
        self.result.counts["candidates added"] += len(new_people)
        self.save_history(plan, sheets)

    def fill_person(self, person: Candidate, group):
        def put(field_name, value):
            """Rows are in order of reliability, so the first value found is kept."""
            if value not in (None, "", []) and getattr(person, field_name) in (None, "", []):
                setattr(person, field_name, value)

        names = [r for r in group if r.name]
        if names:
            first, last = split_name(names[0].name)
            put("first_name", first)
            put("last_name", last)
        emails = [e for r in group for e in r.emails if e not in self.placeholders]
        phones = [p for r in group for p in r.phones if p not in self.placeholders]
        emails = list(dict.fromkeys(emails))
        phones = list(dict.fromkeys(phones))
        put("email", emails[0][:254] if emails else "")
        put("phone", phones[0] if phones else "")
        known = {person.email.lower(), person.phone}
        extra = [x for x in emails + phones if x not in known]
        if extra:
            merged = [x.strip() for x in person.other_contacts.split(",") if x.strip()] if person.other_contacts else []
            person.other_contacts = ", ".join(dict.fromkeys(merged + extra))[:255]

        records = list(person.source_records or [])
        for r in group:
            row = r.data
            if r.kind == "cv":
                sheet, number = r.ref
                put("gender", gender(row.get("gender")))
                put("date_of_birth", self.birth_date(sheet, number, "Date of Birth", row.get("date of birth")))
                put("state_of_origin", state_name(row.get("state"))[:80])
                put("lga", clean(row.get("lga/community tag"))[:150])
                put("relevant_experience", yes_no(row.get("relevant experience")))
                years = to_number(row.get("experience in years"))
                if clean(row.get("experience in years")) and years is None:
                    self.issue(sheet, number, "Experience in Years", row.get("experience in years"), "Not a number; left blank")
                if years is not None and years <= 50:
                    put("years_experience", years.quantize(Decimal("0.1")))
                level = qualification(row.get("degree"))
                if clean(row.get("degree")) and not level and clean(row.get("degree")).lower() not in {"others", "other"}:
                    self.issue(sheet, number, "Degree", row.get("degree"), "Qualification not recognised; kept in the original record")
                put("highest_qualification", level)
                put("course", clean(row.get("discipline"))[:150])
                put("institution", clean(row.get("name of institution"))[:200])
                year = to_number(row.get("year of graduation"))
                if year and 1950 <= year <= self.today.year + 1:
                    put("graduation_year", int(year))
                elif clean(row.get("year of graduation")):
                    self.issue(sheet, number, "Year of Graduation", row.get("year of graduation"), "Not a valid year; left blank")
                put("degree_class", degree_class(row.get("marks_obtained")))
                source = clean(row.get("source"))
                put("source_detail", SOURCE_SPELLINGS.get(source.lower(), source)[:80])
                key = (sheet.title, number)
                if not any((rec.get("sheet"), rec.get("row")) == key for rec in records):
                    records.append({"file": sheet.file, "sheet": sheet.title, "row": number,
                                    "values": sheet.original(row)})
            elif r.kind == "interview":
                put("gender", gender(row.get("gender")))
                community = clean(row.get("community query")).lower()
                put("community_status", {"community": CommunityStatus.COMMUNITY,
                                         "non community": CommunityStatus.NON_COMMUNITY}.get(community, ""))
        person.source_records = records
        if not person.source_detail:
            kinds = {r.kind for r in group}
            person.source_detail = "Interview records" if "interview" in kinds else "CBT records" if "cbt" in kinds else ""

    # -- CBTs, interview records, applications
    def save_history(self, plan, sheets):
        # Rows already in the portal, by their full content and by the event they describe (person, date, test
        # or interview). A row seen before is skipped; a row whose event is known but whose result or score was
        # changed in Excel updates that record; anything else is new.
        known_cbt, cbt_events = defaultdict(list), defaultdict(list)
        for pk, *key in CBTResult.objects.values_list("pk", "candidate_id", "test_date", "category", "score", "total"):
            known_cbt[tuple(key)].append(pk)
            cbt_events[tuple(key[:3])].append(pk)
        known_iv, iv_events = defaultdict(list), defaultdict(list)
        for pk, *key in InterviewRecord.objects.values_list("pk", "candidate_id", "interview_date", "ref_id",
                                                            "panel_category", "result", "score", "panellists"):
            known_iv[tuple(key)].append(pk)
            iv_events[tuple(key[:4])].append(pk)
        seen_cbt, seen_iv, used = set(), set(), set()
        cbts, records, pairs = [], [], defaultdict(list)
        changed_cbt, changed_iv = [], []
        for person, group in plan:
            for r in group:
                sheet, number = r.ref if r.kind in {"cbt", "interview"} else (None, None)
                row = r.data
                if r.kind == "cbt":
                    test_date = self.event_date(sheet, number, "Date", row.get("date"))
                    score, total = to_number(row.get("exam score")), to_number(row.get("total score"))
                    category = clean(row.get("category"))[:80]
                    key = (person.pk, test_date, category, score, total)
                    if key in seen_cbt:
                        self.result.counts["CBT rows repeated (same person, date, test and score; skipped)"] += 1
                        continue
                    seen_cbt.add(key)
                    match = next((pk for pk in known_cbt.get(key, []) if ("cbt", pk) not in used), None)
                    if match:
                        used.add(("cbt", match))
                        continue
                    cbt = CBTResult(candidate=person, test_date=test_date, category=category, score=score,
                                    total=total, company=clean(row.get("company"))[:20], source_row=number)
                    changed_cbt.append((cbt, (person.pk, test_date, category)))
                elif r.kind == "interview":
                    ref = clean(row.get("ref_id"))
                    result_text = clean(row.get("remarks_query")).lower()
                    result = INTERVIEW_RESULTS.get(result_text, "")
                    if result_text and not result:
                        self.issue(sheet, number, "Remarks_Query", row.get("remarks_query"), "Result not recognised")
                    when = self.event_date(sheet, number, "Date", row.get("date"))
                    score = to_number(row.get("score"))
                    if score is not None and score <= 1:
                        score = (score * 100).quantize(Decimal("0.1"))
                    panel_category = clean(row.get("panel category"))[:80]
                    panellists = clean(row.get("panellists"))[:255]
                    key = (person.pk, when, ref[:30], panel_category, result, score, panellists)
                    if key in seen_iv:
                        self.result.counts["interview rows repeated (same person, date, panel and result; skipped)"] += 1
                        continue
                    seen_iv.add(key)
                    match = next((pk for pk in known_iv.get(key, []) if ("iv", pk) not in used), None)
                    if match:
                        used.add(("iv", match))
                        continue
                    requisition = self.requisitions.get(ref.lower()) if ref else None
                    if ref and not requisition:
                        self.result.counts["interview rows whose requisition is not in the tracker"] += 1
                    members = " / ".join(x for x in (clean(row.get("panel member")), clean(row.get("panel member2"))) if x)
                    record = InterviewRecord(
                        candidate=person, requisition=requisition, ref_id=ref[:30], interview_date=when,
                        classification=clean(row.get("classification"))[:40], panel_category=panel_category,
                        panellists=panellists, panel_members=members[:255], score=score,
                        result=result, community_status=clean(row.get("community query"))[:20],
                        onboarding_status=clean(row.get("onboarding_status"))[:40],
                        join_date=self.event_date(sheet, number, "Join Date", row.get("join date")),
                        remark=clean(row.get("remark"))[:60], source_row=number,
                    )
                    changed_iv.append((record, (person.pk, when, ref[:30], panel_category)))

        # Rows that are not exactly as before: an edit of a known event, or new.
        cbt_updates, iv_updates = [], []
        for cbt, event in changed_cbt:
            pk = next((pk for pk in cbt_events.get(event, []) if ("cbt", pk) not in used), None)
            if pk:
                used.add(("cbt", pk))
                cbt.pk = pk
                cbt_updates.append(cbt)
            else:
                cbts.append(cbt)
        stored = InterviewRecord.objects.in_bulk(
            [pk for _, event in changed_iv for pk in iv_events.get(event, [])])
        for record, event in changed_iv:
            pk = next((pk for pk in iv_events.get(event, []) if ("iv", pk) not in used), None)
            if pk:
                used.add(("iv", pk))
                record.pk, record.application_id = pk, stored[pk].application_id
                iv_updates.append(record)
            else:
                records.append(record)
            if record.requisition:
                pairs[(record.candidate, record.requisition)].append(record)

        self.log(f"Saving {len(cbts):,} CBT results and {len(records):,} interview records...")
        CBTResult.objects.bulk_create(cbts, batch_size=1000)
        CBTResult.objects.bulk_update(cbt_updates, ["score", "total", "company", "source_row"], batch_size=1000)
        self.result.counts["CBT results"] += len(cbts)
        if cbt_updates:
            self.result.counts["CBT results changed in Excel (updated)"] += len(cbt_updates)
        self.save_applications(pairs)
        InterviewRecord.objects.bulk_create(records, batch_size=1000)
        InterviewRecord.objects.bulk_update(iv_updates, [
            "requisition", "application", "classification", "panellists", "panel_members", "score", "result",
            "community_status", "onboarding_status", "join_date", "remark", "source_row"], batch_size=1000)
        self.result.counts["interview records"] += len(records)
        if iv_updates:
            self.result.counts["interview records changed in Excel (updated)"] += len(iv_updates)

    PRIORITY = {"onboarded": 7, "in progress": 6, "selected": 5, "on_hold": 4, "rejected": 3, "declined": 2,
                "not_found": 1, "rescheduled": 0}

    def outcome(self, record: InterviewRecord) -> str:
        onboarding = record.onboarding_status.lower()
        if onboarding.startswith("onboarded"):
            return "onboarded"
        if onboarding == "in progress":
            return "in progress"
        return record.result or "rescheduled"

    def save_applications(self, pairs):
        existing = set(Application.objects.values_list("candidate_id", "requisition_id"))
        for (person, requisition), records in pairs.items():
            if (person.pk, requisition.pk) in existing:
                app = Application.objects.get(candidate=person, requisition=requisition)
                for record in records:
                    record.application = app
                # Untouched in the portal: the application follows what Excel now says about the interviews.
                if not self.replace and app.source == Candidate.Source.IMPORT \
                        and not Activity.objects.filter(application=app).exists():
                    history = {r.pk: r for r in app.interview_records.all()}
                    history.update({r.pk: r for r in records if r.pk})
                    before = (app.stage, app.status, app.decision)
                    outcome = self.apply_outcome(app, list(history.values()) + [r for r in records if not r.pk],
                                                 requisition)
                    if (app.stage, app.status, app.decision) != before:
                        app.save()
                        self.start_onboarding(app, outcome)
                        self.result.counts["applications moved on by the interview records"] += 1
                continue
            app = Application(candidate=person, requisition=requisition, source=Candidate.Source.IMPORT,
                              match_score=Decimal("0"), match_grade="")
            first = min((r.interview_date for r in records if r.interview_date), default=None)
            app.applied_at = aware(first) or aware(requisition.submitted_at.date() if requisition.submitted_at else None) \
                or timezone.now()
            outcome = self.apply_outcome(app, records, requisition)
            app.save()
            self.start_onboarding(app, outcome)
            for record in records:
                record.application = app
            self.result.counts["applications"] += 1
            self.result.counts[f"applications: {app.get_status_display().lower()}"] += 1

    def apply_outcome(self, app, records, requisition) -> str:
        """Stage, status and decision from the latest interview record. Returns the outcome."""
        records.sort(key=lambda r: ((r.interview_date or date.min), self.PRIORITY[self.outcome(r)]))
        final = records[-1]
        outcome = self.outcome(final)
        open_req = requisition.status in Requisition.ACTIVE_STATUSES
        first = min((r.interview_date for r in records if r.interview_date), default=None)
        last = final.interview_date or first
        app.stage_changed_at = aware(last) or app.applied_at
        app.decision, app.status_reason, app.hired_at = "", "", None
        recorded = final.get_result_display() if final.result else "no result"
        if outcome == "onboarded":
            app.stage, app.status, app.decision = Stage.HIRED, Application.Status.HIRED, Application.Decision.SELECT
            app.hired_at = aware(final.join_date) or app.stage_changed_at
            app.status_reason = f"Onboarded ({final.onboarding_status.split('_')[-1]})"
        elif outcome == "in progress":
            app.stage, app.status, app.decision = Stage.ONBOARDING, Application.Status.ACTIVE, Application.Decision.SELECT
        elif outcome == "selected":
            app.stage, app.decision = Stage.DECISION, Application.Decision.SELECT
            app.status = Application.Status.ACTIVE if open_req else Application.Status.NOT_ENGAGED
        elif outcome == "on_hold":
            app.stage, app.decision = Stage.DECISION, Application.Decision.ON_HOLD
            app.status = Application.Status.ON_HOLD if open_req else Application.Status.NOT_ENGAGED
        elif outcome == "rejected":
            app.stage, app.status, app.decision = Stage.DECISION, Application.Status.REJECTED, Application.Decision.REJECT
        elif outcome == "declined":
            app.stage, app.status = Stage.DECISION, Application.Status.WITHDRAWN
            app.status_reason = "Declined (as recorded)"
        else:  # rescheduled, not found
            app.stage = Stage.INTERVIEW
            app.status = Application.Status.ACTIVE if open_req and outcome == "rescheduled" else Application.Status.NOT_ENGAGED
        if app.status == Application.Status.NOT_ENGAGED:
            app.status_reason = f"Not engaged; last recorded result: {recorded}"
        app.decided_at = app.stage_changed_at if app.decision else None
        app._join_date = final.join_date
        return outcome

    def start_onboarding(self, app, outcome):
        if outcome != "in progress" or Onboarding.objects.filter(application=app).exists():
            return
        onboarding = Onboarding.objects.create(application=app, start_date=app._join_date,
                                               status=Onboarding.Status.IN_PROGRESS,
                                               notes="Onboarding in progress when imported from the interview records.")
        OnboardingTask.objects.bulk_create([OnboardingTask(onboarding=onboarding, title=t, order=i)
                                            for i, t in enumerate(DEFAULT_ONBOARDING_TASKS)])
        self.result.counts["onboardings in progress"] += 1


# --------------------------------------------------------------------------- report
def write_report(result: ImportResult, folder) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"import-report-{timezone.localtime():%Y-%m-%d-%H%M%S}.xlsx")
    wb = Workbook()
    bold = Font(name="Arial", bold=True)
    plain = Font(name="Arial")

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        for col, width in zip("ABCDEFG", widths):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows():
            for cell in row:
                cell.font = bold if cell.row == 1 else plain
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Summary"
    summary = [("Imported", k, v) for k, v in sorted(result.counts.items())]
    summary += [("Removed before import", k, v) for k, v in sorted(result.removed.items()) if v]
    summary += [("Placeholder contact (not used to match people)", k, f"{n} rows")
                for k, n in result.placeholders.items()]
    sheet(ws, ["What", "Item", "Number"], summary, [44, 52, 14])
    sheet(wb.create_sheet("Values not used"), ["File", "Sheet", "Row", "Column", "Value", "What happened"],
          result.issues, [36, 30, 8, 30, 30, 60])
    sheet(wb.create_sheet("Shared contacts"), ["Email or phone", "First person", "Other person"],
          result.shared_contacts, [32, 50, 50])
    sheet(wb.create_sheet("Same name, different contacts"), ["Name", "People with this name (email, phone, source)"],
          result.same_names, [36, 120])
    sheet(wb.create_sheet("Duplicates merged"), ["Candidate", "Rows merged", "From"],
          result.merged, [40, 14, 50])
    wb.save(path)
    return path
